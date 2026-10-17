"""DOCX 파서. 본문(w:body)을 처음부터 순서대로 읽는다.

- 제목: Heading 1~3 스타일 → level 1~3, 표지의 큰 글자(20pt 이상) 또는 Title 스타일 → level 0
- '문서 정보'·'Contents' 쪽 제목(굵은 일반 문단)은 제목으로 넘겨 공통 후처리에서 정리한다
- 목록: List 스타일·번호 매기기, 또는 '•'·'①'로 시작하는 문단
- 표: 2줄 카드 표 → 카드 제목(level 3) + 목록, 1칸 표(안내·인용·코드 상자) → 안의 문단, 그 밖 → 표
- 그림: 문단 안 그림 → image, 바로 뒤 '그림 …' 문단을 설명으로 붙인다
"""
from __future__ import annotations

import os
import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from .common import CAPTION_RE, SKIP_TITLES, clean, finalize, split_bullet
from .models import Element, ParsedDocument

R_EMBED = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'


def _style(p: Paragraph) -> str:
    try:
        return (p.style.name or '').lower() if p.style is not None else ''
    except (KeyError, AttributeError):
        return ''


def _images(el, part) -> list[str]:
    names = []
    for blip in el.xpath('.//*[local-name()="blip"]'):
        rid = blip.get(R_EMBED)
        if rid and rid in part.related_parts:
            names.append(os.path.basename(str(part.related_parts[rid].partname)))
    return names


def _cell_text(cell) -> str:
    """칸의 글자. 칸 안에 표가 또 있으면 그 글자도 넣는다 (python-docx 의 cell.text 는 안쪽 표를 빠뜨린다)."""
    parts = []
    for child in cell._tc.iterchildren():
        tag = child.tag.split('}')[1]
        if tag == 'p':
            parts.append(Paragraph(child, cell).text)
        elif tag == 'tbl':
            parts += [' | '.join(r) for r in _rows(Table(child, cell))]
    return clean('\n'.join(parts))


def _rows(t: Table) -> list[list[str]]:
    rows = []
    for r in t.rows:
        cells, prev = [], None
        for c in r.cells:
            if c._tc is prev:          # 가로 병합 칸은 한 번만
                continue
            prev = c._tc
            cells.append(_cell_text(c))
        if any(cells):
            rows.append(cells)
    return rows


def _is_cards(rows: list[list[str]]) -> bool:
    """카드 표: 2줄, 첫 줄은 짧은 제목, 둘째 줄은 칸마다 '·'로 시작하는 목록."""
    if len(rows) != 2 or len(rows[0]) != len(rows[1]) or not rows[0]:
        return False
    if any('\n' in c or not c or len(c) > 40 for c in rows[0]):
        return False
    return any(rows[1]) and all(c.startswith(('·', '•')) for c in rows[1] if c)


class _Builder:
    def __init__(self, d):
        self.d = d
        self.out: list[Element] = []
        self.pending_image: Element | None = None
        self.title_done = False

    def add(self, kind: str, text: str = '', **kw) -> Element:
        e = Element(kind=kind, text=text, **kw)
        self.out.append(e)
        return e

    # ---- 문단
    def paragraph(self, p: Paragraph, in_cell: bool = False):
        for name in _images(p._p, self.d.part):
            self.pending_image = self.add('image', meta={'name': name})
        text = clean((p.text or '').replace('\t', ' '))
        if not text:
            return
        style = _style(p)
        if style.startswith(('heading', 'title')):
            self._heading(' '.join(text.split('\n')), style)
            return
        runs = [r for r in p.runs if r.text.strip()]
        bold = bool(runs) and all(r.bold for r in runs)
        size = max((r.font.size.pt for r in runs if r.font.size), default=0)
        is_list = style.startswith('list') or bool(p._p.xpath('./w:pPr/w:numPr'))
        for line in text.split('\n'):
            self._line(line, bold, size, is_list, in_cell)

    def _heading(self, text: str, style: str):
        self.pending_image = None
        if style == 'title':
            if not self.title_done:
                self.title_done = True
                self.add('heading', text, level=0)
            return
        m = re.match(r'heading (\d)', style)
        self.add('heading', text, level=min(max(int(m.group(1)), 1), 3) if m else 1)

    def _line(self, text: str, bold: bool, size: float, is_list: bool, in_cell: bool):
        if self.pending_image is not None and CAPTION_RE.match(text):
            self.pending_image.text = text          # 그림 설명
            self.pending_image = None
            return
        self.pending_image = None
        if not in_cell and bold and text in SKIP_TITLES:
            self.add('heading', text, level=1)      # '문서 정보'·'Contents' 쪽 제목 → 공통 후처리에서 정리
            return
        if not in_cell and not self.title_done and size >= 20:
            self.title_done = True                  # 표지의 큰 제목 (Title 스타일이 없는 문서)
            self.add('heading', text, level=0)
            return
        bullet, body = split_bullet(text)
        if is_list or bullet:
            self.add('list_item', body)
        else:
            self.add('paragraph', text)

    # ---- 표
    def table(self, t: Table):
        rows = _rows(t)
        images = _images(t._tbl, self.d.part)
        if not rows:
            for name in images:
                self.pending_image = self.add('image', meta={'name': name})
            return
        if _is_cards(rows):
            for title, body in zip(rows[0], rows[1]):
                self.add('heading', title, level=3, meta={'card': True})
                for line in body.split('\n'):
                    if line.strip():
                        self.add('list_item', split_bullet(line.strip())[1])
            return
        if len(t.rows) == 1 and len(t.columns) == 1:
            # 안내·인용·코드 상자: 표가 아니라 안의 문단
            cell = t.cell(0, 0)
            for child in cell._tc.iterchildren():
                tag = child.tag.split('}')[1]
                if tag == 'p':
                    self.paragraph(Paragraph(child, cell), in_cell=True)
                elif tag == 'tbl':
                    self.table(Table(child, cell))
            return
        for name in images:
            self.add('image', meta={'name': name})
        self.add('table', rows=rows)


def parse_docx(path: str, safe: bool = False) -> ParsedDocument:
    d = Document(path)
    cp = d.core_properties
    doc = ParsedDocument(path=path, format='docx',
                         properties={'title': cp.title or '', 'author': cp.author or '',
                                     'modified': str(cp.modified or ''), 'keywords': cp.keywords or ''})
    b = _Builder(d)
    for el in d.element.body.iterchildren():
        tag = el.tag.split('}')[1]
        if tag == 'p':
            b.paragraph(Paragraph(el, d))
        elif tag == 'tbl':
            b.table(Table(el, d))
    doc.elements = b.out
    return finalize(doc, safe)
