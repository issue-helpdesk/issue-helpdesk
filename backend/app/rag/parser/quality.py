"""누락 감지. 파일의 글자를 형식별로 '있는 그대로' 다시 뽑아, 파싱 결과에 남았는지 확인한다.

파서와 일부러 다른 길로 원문을 읽는다(docx 모든 문단, pptx 모든 도형·표, PDF 모든 줄, md·txt 모든 줄).
같은 코드로 읽으면 파서의 실수를 검사도 똑같이 놓치기 때문이다.

결과에 '남은' 곳: 본문 요소, 표 칸, 항목 머리(props), 문서 정보(doc.info), 템플릿 표지 문구(300자 이하),
정당한 이유로 뺀 글자(목차·쪽 제목·장 번호·쪽 번호·반복 제목, 두 쪽 이상 반복되는 꼬리말).

검사 항목
- coverage        원문 조각 글자 중 결과에 있는 비율. 조각은 결과 요소 하나(또는 이웃한 요소 6개 이내)에 들어 있어야 하고,
                  원문에 같은 조각이 여러 번 나오면 결과에도 그 횟수만큼 있어야 한다
- long_missing    빠진 조각 중 40자 이상 (긴 문서에서 비율로는 가려지는 문단 단위 누락)
- missing_values  숫자·코드(IT-SW-002, FAQ-023, AUTH-101, 443, 2026-10-16 …)의 빠진 개수
- dropped_ratio   정당한 제외로 뺀 서로 다른 글자의 비율. 15%를 넘으면 배치 규칙이 본문을 빼는 것으로 의심한다
"""
from __future__ import annotations

import bisect
import re
from collections import Counter

from .common import CIRCLED, TOC_TITLES, clean
from .models import ParsedDocument

MIN_COVERAGE = 0.98
LONG_PIECE = 40                    # 이 길이(정규화 글자) 이상 조각이 빠지면 비율과 상관없이 문제
MAX_DROPPED_RATIO = 0.15
COVER_MAX_CHARS = 300              # 이보다 긴 '표지'는 본문이 섞인 첫 쪽으로 보고 남은 글자로 치지 않는다
WINDOW = 6                         # 조각 하나가 걸칠 수 있는 이웃 요소 수 (줄이 여러 요소로 나뉜 경우)
ACCEPTED_DROPS = ('toc', 'skip_title', 'divider', 'page_no', 'repeat_title')

_STRIP_RE = re.compile(r'[\s•·▪◦‣\-–*#|>`!\[\]()]+')
_VALUE_RE = re.compile(r'[A-Za-z]{2,}(?:-[A-Za-z0-9]+)*-\d+[A-Za-z0-9-]*'      # IT-SW-002, FAQ-023, AUTH-101
                       r'|\d+(?:[.:/-]\d+)+'                                   # 2026-10-16, 10.0.0.1, 09:00
                       r'|\d{2,}')                                             # 443, 1000


def norm(text: str) -> str:
    """비교용: '(계속)'·공백·글머리·마크다운 기호를 지우고 ①을 1. 로."""
    text = text.replace('(계속)', '')
    text = re.sub(f'[{CIRCLED}]', lambda m: f'{CIRCLED.index(m.group()) + 1}.', text)
    return _STRIP_RE.sub('', text)


# ---------------------------------------------------------------- 결과 쪽: 남은 글자

def kept_units(doc: ParsedDocument) -> list[str]:
    """결과에 남은 글자를 요소 단위로 (읽는 순서)."""
    units = []
    for e in doc.elements:
        if e.text:
            units.append(e.text)
        units += [c for r in (e.rows or []) for c in r if c]
        units += [f'{k} {v}' for k, v in e.meta.get('props', []) or []]
    units += [f'{k} {v}' for k, v in doc.info.items()]
    cover = doc.properties.get('cover', [])
    if sum(len(c) for c in cover) <= COVER_MAX_CHARS:
        units += cover
    units += accepted_drops(doc.dropped)
    return units


def accepted_drops(dropped: list) -> list[str]:
    """정당한 제외로 인정하는 뺀 글자. 꼬리말·제목 띠 글자는 두 쪽 이상에 반복될 때만 인정한다
    (한 번만 나오는 띠 글자는 배치가 다른 문서의 본문일 수 있다)."""
    key = lambda t: re.sub(r'\d+', '#', t)          # noqa: E731
    pages: dict[str, set] = {}
    for why, text, page in dropped:
        if why == 'band':
            pages.setdefault(key(text), set()).add(page)
    return [text for why, text, page in dropped
            if why in ACCEPTED_DROPS or (why == 'band' and len(pages.get(key(text), ())) >= 2)]


class _Index:
    """남은 글자를 이어 붙이고 요소 경계를 기억한다. 조각이 경계를 WINDOW-1 번 넘지 않고 들어 있는 횟수를 센다."""

    def __init__(self, units: list[str]):
        normed = [norm(u) for u in units]
        self.text = ''.join(normed)
        self.bounds = []
        pos = 0
        for n in normed:
            pos += len(n)
            self.bounds.append(pos)

    def count(self, piece: str) -> int:
        n, start = 0, 0
        while (i := self.text.find(piece, start)) >= 0:
            crossed = bisect.bisect_left(self.bounds, i + len(piece)) - bisect.bisect_right(self.bounds, i)
            if crossed < WINDOW:
                n += 1
                start = i + len(piece)        # 겹치지 않게 센다
            else:
                start = i + 1
        return n


# ---------------------------------------------------------------- 원문 쪽: 있는 그대로 다시 읽기

def raw_pieces(path: str, fmt: str) -> list[str]:
    """형식별 원문 조각 (문단·도형 줄·표 칸·PDF 줄 단위)."""
    if fmt == 'docx':
        from docx import Document
        ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        out = []
        for p in Document(path).element.body.iter(ns + 'p'):
            parts = []
            for el in p.iter(ns + 't', ns + 'br', ns + 'tab'):
                parts.append(el.text or '' if el.tag == ns + 't' else ('\n' if el.tag == ns + 'br' else ' '))
            out += ''.join(parts).split('\n')          # 줄바꿈(<w:br/>)마다 한 조각
        return out
    if fmt == 'pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        out = []

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(sh.shapes)
                elif getattr(sh, 'has_table', False) and sh.has_table:
                    out.extend(c.text for r in sh.table.rows for c in r.cells)
                elif getattr(sh, 'has_text_frame', False) and sh.has_text_frame:
                    out.extend(p.text for p in sh.text_frame.paragraphs)
        for s in Presentation(path).slides:
            walk(s.shapes)
        return out
    if fmt == 'pdf':
        import pdfplumber

        from .pdf_parser import _segments
        out = []
        with pdfplumber.open(path) as pdf:
            for pg in pdf.pages:          # 나란한 칸이 한 줄로 합쳐지지 않게 줄 나누기만 파서와 같은 방식을 쓴다
                out.extend(s.text for s in _segments(pg.chars))
        return out
    if fmt in ('md', 'txt'):
        return _md_pieces(path)
    return []


def _md_pieces(path: str) -> list[str]:
    """md·txt 원문 줄에서 마크다운 기호만 걷어 낸 조각. 머리말과 '목차' 절은 뺀다."""
    from .md_parser import FRONT_RE, IMAGE_RE, _inline, read_text
    text, _ = read_text(path)
    m = FRONT_RE.match(text)
    out, in_toc, alt = [], False, None
    for line in (text[m.end():] if m else text).splitlines():
        s = line.strip()
        if h := re.match(r'^#{1,6}\s+(.*)$', s):
            in_toc = _inline(h.group(1)) in TOC_TITLES
            s = h.group(1)
        if in_toc or not s or s.startswith('```') or re.fullmatch(r'\|?[\s:|-]+\|?', s):
            continue
        if alt is not None and re.fullmatch(r'\*[^*].*\*', s) and _inline(s[1:-1]) == alt:
            alt = None
            continue          # '![설명](…)' 바로 뒤에 같은 설명을 기울임으로 한 번 더 쓴 줄
        alt = None
        if img := IMAGE_RE.match(s):
            s = alt = _inline(img.group(1))
        if s.startswith('|'):
            out += [_inline(c) for c in re.split(r'(?<!\\)\|', s.strip('|'))]
            continue
        s = re.sub(r'^(>\s*)+|^[-*+]\s+|^\d{1,2}\.\s+', '', s)
        out.append(_inline(re.sub(r'^\*([^*].*)\*$', r'\1', s)))
    return out


# ---------------------------------------------------------------- 판정

def assess(doc: ParsedDocument, pieces: list[str]) -> dict:
    units = kept_units(doc)
    index = _Index(units)
    # 장 번호 표기만 결과와 맞춘다: '01  개요' · '01. 개요' → '1. 개요'
    pieces = [re.sub(r'^0(\d)\.?\s+', r'\1. ', clean(p)) for p in pieces]
    groups: dict[str, list[str]] = {}
    for p in pieces:
        if n := norm(p):
            groups.setdefault(n, []).append(p)
    total = lost = 0
    missing: list[str] = []
    for n, originals in groups.items():
        total += len(n) * len(originals)
        short = len(originals) - min(len(originals), index.count(n))
        if short:
            lost += len(n) * short
            missing += originals[:short]
    coverage = 1 - lost / total if total else 1.0
    long_missing = [m for m in missing if len(norm(m)) >= LONG_PIECE]
    values = _missing_values(pieces, units)
    drop_chars = sum(len(norm(t)) for t in set(accepted_drops(doc.dropped)))
    dropped_ratio = drop_chars / total if total else 0.0
    return {
        'coverage': round(coverage, 4), 'missing': missing[:20], 'long_missing': long_missing[:10],
        'missing_values': values[:20], 'dropped_ratio': round(dropped_ratio, 4),
        'ok': coverage >= MIN_COVERAGE and not long_missing and not values and dropped_ratio <= MAX_DROPPED_RATIO,
    }


def _missing_values(pieces: list[str], units: list[str]) -> list[str]:
    """원문에 나온 숫자·코드 중 결과에 한 번도 없는 값. (같은 문장이 여러 번 나오는 횟수는 조각 단위에서 이미 센다.
    값까지 횟수로 세면 머리 띠에 한 번 더 찍힌 문서 번호처럼 정보가 이미 남아 있는 경우도 누락이 된다)"""
    raw = Counter(v for p in pieces for v in _VALUE_RE.findall(p))
    kept = {v for u in units for v in _VALUE_RE.findall(u)}
    return [v for v in raw if v not in kept]


def problems(q: dict) -> list[str]:
    """사람이 읽는 경고 문구."""
    out = []
    if q['coverage'] < MIN_COVERAGE:
        out.append(f"원문 글자 포함률 {q['coverage']:.1%}")
    if q['long_missing']:
        out.append(f"빠진 문단 {len(q['long_missing'])}개 (예: {q['long_missing'][0][:40]}…)")
    if q['missing_values']:
        out.append(f"빠진 숫자·코드 {len(q['missing_values'])}개 (예: {', '.join(q['missing_values'][:5])})")
    if q['dropped_ratio'] > MAX_DROPPED_RATIO:
        out.append(f"정상 제외로 뺀 글자가 {q['dropped_ratio']:.0%} — 배치 규칙이 본문을 뺐을 수 있음")
    return out
