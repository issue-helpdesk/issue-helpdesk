"""가이드·매뉴얼 docx 빌더 (16:9 가로, 장표형).

content_*.py 의 DOC 사전을 받아 docx를 만든다.
목차의 쪽 번호는 PDF로 한 번 변환한 뒤 장 시작 쪽을 찾아 다시 채운다 (make_all.py).
"""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

FONT = '맑은 고딕'
NAVY = '22395F'
NAVY_LIGHT = 'E9EEF4'
ORANGE = 'C0461A'
ACCENT = 'E8590C'
GREY = '5B6573'
LINE = 'D5DBE3'
WARN_BG = 'FBE8E3'
WARN_LINE = 'E7B3A6'
NOTE_BG = 'F3F5F8'
OK_BG = 'E8F3EC'

PAGE_W, PAGE_H = Cm(33.867), Cm(19.05)
MARGIN_LR, MARGIN_T, MARGIN_B = Cm(1.8), Cm(1.1), Cm(1.3)
BODY_W = PAGE_W - 2 * MARGIN_LR

# 처리 구분 (요청 표의 🔴🟡🟢를 글자 배지로 표시)
LEVEL = {
    'red': ('즉시 신고·이관', 'C0392B', 'FBE3E0'),
    'yellow': ('조건부 이관·승인 필요', 'B7791F', 'FCF1DC'),
    'green': ('셀프 조치 (AI 안내)', '2E7D4F', 'E3F2E8'),
}


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _cell_borders(cell, **edges):
    """edges: top/left/bottom/right = (size_eighths, color) 또는 None(없음)."""
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        v = edges.get(edge)
        if v is None:
            el.set(qn('w:val'), 'nil')
        else:
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), str(v[0]))
            el.set(qn('w:color'), v[1])
        b.append(el)
    tcPr.append(b)


def _cell_margins(cell, top=80, bottom=80, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        el = OxmlElement(f'w:{k}')
        el.set(qn('w:w'), str(v))
        el.set(qn('w:type'), 'dxa')
        m.append(el)
    tcPr.append(m)


def _no_table_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        b.append(el)
    tblPr.append(b)


def _fixed_layout(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[i] / 635)))


def _keep_row(row):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement('w:cantSplit')
    trPr.append(cs)


def _border_bottom(paragraph, color=NAVY, size=12, space=4):
    pPr = paragraph._p.get_or_add_pPr()
    pb = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'), str(size))
    bot.set(qn('w:space'), str(space))
    bot.set(qn('w:color'), color)
    pb.append(bot)
    pPr.append(pb)


def _field(run, instr):
    for kind, text in (('begin', None), ('instr', instr), ('separate', None), ('text', '1'), ('end', None)):
        if kind == 'instr':
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = f' {instr} '
        elif kind == 'text':
            el = OxmlElement('w:t')
            el.text = text
        else:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), kind)
        run._r.append(el)


def set_run(run, size=None, bold=None, color=None, font=FONT, mono=False):
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = rgb(color)
    name = 'Consolas' if mono else font
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.append(rf)
    rf.set(qn('w:eastAsia'), FONT)
    rf.set(qn('w:ascii'), name)
    rf.set(qn('w:hAnsi'), name)


def para_fmt(p, before=0, after=2, line=1.15, align=None, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align:
        p.alignment = align
    if keep:
        pf.keep_with_next = True


def rich(p, text, size=11, color='1F2933', bold=False):
    """'**굵게**', '`코드`' 표기를 지원하는 한 줄 출력."""
    import re
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**'):
            r = p.add_run(part[2:-2]); set_run(r, size, True, color)
        elif part.startswith('`'):
            r = p.add_run(part[1:-1]); set_run(r, size - 0.5, False, NAVY, mono=True)
        else:
            r = p.add_run(part); set_run(r, size, bold, color)
    return p


class Builder:
    def __init__(self, meta):
        self.meta = meta
        self.doc = Document()
        self._setup()
        self.first_page = True
        self.toc_pages = {}

    # ---------- 문서 설정 ----------
    def _setup(self):
        s = self.doc.sections[0]
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = PAGE_W, PAGE_H
        s.left_margin = s.right_margin = MARGIN_LR
        s.top_margin, s.bottom_margin = MARGIN_T, MARGIN_B
        s.footer_distance = Cm(0.5)
        st = self.doc.styles['Normal']
        st.font.name = FONT
        st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        for name, size, color in (('Heading 1', 26, NAVY), ('Heading 2', 15, ORANGE), ('Title', 34, '1F2933')):
            hs = self.doc.styles[name]
            hs.font.name = FONT
            hs.font.size = Pt(size)
            hs.font.bold = True
            hs.font.color.rgb = rgb(color)
            hs.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
            hs.element.rPr.rFonts.set(qn('w:ascii'), FONT)
            hs.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
            pf = hs.paragraph_format
            pf.space_before, pf.space_after = Pt(0), Pt(4)
            # Title 스타일의 아래 테두리 제거
            pPr = hs.element.find(qn('w:pPr'))
            if pPr is not None:
                for b in pPr.findall(qn('w:pBdr')):
                    pPr.remove(b)
        # 바닥글: 왼쪽 저작 표시, 오른쪽 쪽 번호
        ft = s.footer.add_table(1, 2, BODY_W)
        _no_table_borders(ft)
        _fixed_layout(ft, [BODY_W - Cm(4), Cm(4)])
        for cell in ft.rows[0].cells:
            _cell_margins(cell, 0, 0, 0, 0)
        f = ft.rows[0].cells[0].paragraphs[0]
        r = f.add_run(self.meta.get('footer', '© HanwhaSystems (가칭) · 교육용 가상 문서')); set_run(r, 9, False, GREY)
        tail = s.footer.paragraphs[-1]
        para_fmt(tail, 0, 0, 1.0)
        tr = tail.add_run(''); set_run(tr, 1)
        tail.paragraph_format.line_spacing = Pt(1)
        f = ft.rows[0].cells[1].paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = f.add_run(); set_run(r, 9, False, GREY); _field(r, 'PAGE')
        r = f.add_run(' / '); set_run(r, 9, False, GREY)
        r = f.add_run(); set_run(r, 9, False, GREY); _field(r, 'NUMPAGES')
        cp = self.doc.core_properties
        cp.title = f"{self.meta['doc_no']} {self.meta['title']}"
        cp.author = self.meta['owner']
        cp.last_modified_by = self.meta['owner']
        cp.subject = self.meta.get('category', '')
        cp.keywords = self.meta.get('keywords', '')
        cp.comments = '교육 프로젝트용 가상 문서'

    # ---------- 페이지 ----------
    def new_page(self):
        if self.first_page:
            self.first_page = False
            return None
        p = self.doc.add_paragraph()
        para_fmt(p, 0, 0, 1.0)
        p.add_run().add_break(WD_BREAK.PAGE)
        return p

    def _first_para(self):
        """새 페이지의 첫 문단. 페이지 나눔 문단을 따로 두지 않도록 page_break_before 사용."""
        p = self.doc.add_paragraph()
        if self.first_page:
            self.first_page = False
        else:
            p.paragraph_format.page_break_before = True
        return p

    def page_header(self, title, crumb, note=None):
        p = self._first_para()
        para_fmt(p, 0, 6, 1.0)
        p.paragraph_format.tab_stops.add_tab_stop(BODY_W, WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(title); set_run(r, 22, True, '1F2933')
        r = p.add_run('\t' + crumb); set_run(r, 10.5, True, NAVY)
        if note:
            r = p.add_run(); r.add_break()
            r = p.add_run('\t' + note); set_run(r, 9.5, False, GREY)
        _border_bottom(p, NAVY, 14, 6)
        sp = self.doc.add_paragraph(); para_fmt(sp, 0, 2, 0.6)

    def cover(self):
        m = self.meta
        p = self._first_para(); para_fmt(p, 70, 0, 1.0)
        r = p.add_run('━━━━━'); set_run(r, 16, True, ACCENT)
        p = self.doc.add_paragraph(); para_fmt(p, 18, 2, 1.0)
        r = p.add_run(m['series']); set_run(r, 14, False, GREY)
        p = self.doc.add_paragraph(style='Title'); para_fmt(p, 2, 6, 1.0)
        r = p.add_run(m['title']); set_run(r, 34, True, '1F2933')
        p = self.doc.add_paragraph(); para_fmt(p, 0, 0, 1.0)
        r = p.add_run(f"{m['doc_no']}  |  {m['owner']}  |  {m['date_dot']}  |  {m['version_label']}")
        set_run(r, 13, True, NAVY)
        p = self.doc.add_paragraph(); para_fmt(p, 95, 0, 1.0)
        r = p.add_run(m.get('cover_note', '본 문서는 교육 프로젝트용 가상 문서입니다. 실제 회사의 공식 정책·절차가 아닙니다.'))
        set_run(r, 10, False, GREY)
        if m.get('confidential'):
            p = self.doc.add_paragraph(); para_fmt(p, 2, 0, 1.0)
            r = p.add_run(m['confidential']); set_run(r, 10, True, 'C0392B')

    def info_page(self, rows, note=None):
        self.page_header('문서 정보', self.meta['doc_no'])
        half = (len(rows) + 1) // 2
        left, right = rows[:half], rows[half:]
        t = self.doc.add_table(rows=max(len(left), len(right)), cols=4)
        _no_table_borders(t)
        w = [Cm(4.2), BODY_W / 2 - Cm(4.6), Cm(4.2), BODY_W / 2 - Cm(3.8)]
        _fixed_layout(t, w)
        for i in range(len(t.rows)):
            for side, data in ((0, left), (2, right)):
                kc, vc = t.rows[i].cells[side], t.rows[i].cells[side + 1]
                for c in (kc, vc):
                    _cell_borders(c, bottom=(4, LINE), top=None, left=None, right=None)
                    _cell_margins(c, 70, 70, 120, 120)
                if i < len(data):
                    k, v = data[i]
                    _shade(kc, NAVY_LIGHT)
                    rich(kc.paragraphs[0], k, 10.5, NAVY, True)
                    rich(vc.paragraphs[0], v, 10.5)
                    for c in (kc, vc):
                        para_fmt(c.paragraphs[0], 0, 0, 1.1)
        if note:
            p = self.doc.add_paragraph(); para_fmt(p, 8, 0)
            rich(p, note, 10, GREY)

    def contents(self, items):
        self.page_header('Contents', '')
        p = self.doc.add_paragraph(); para_fmt(p, 10, 0)
        for label, key in items:
            p = self.doc.add_paragraph(); para_fmt(p, 0, 9, 1.0)
            p.paragraph_format.left_indent = Cm(9)
            p.paragraph_format.tab_stops.add_tab_stop(BODY_W - Cm(2), WD_TAB_ALIGNMENT.RIGHT, 2)  # 점선 leader
            sub = label.startswith('  ')
            r = p.add_run(('    - ' + label.strip()) if sub else label); set_run(r, 11 if sub else 12.5, not sub, '1F2933')
            pg = self.toc_pages.get(key)
            r = p.add_run('\t' + (f'{pg}p' if pg else '')); set_run(r, 12, False, '1F2933')

    def divider(self, num, title, key=None):
        p = self._first_para(); para_fmt(p, 120, 0, 1.0)
        p.paragraph_format.left_indent = Cm(1)
        r = p.add_run(num); set_run(r, 40, True, ACCENT)
        p = self.doc.add_paragraph(style='Heading 1'); para_fmt(p, 0, 0, 1.0)
        p.paragraph_format.left_indent = Cm(1)
        r = p.add_run(title); set_run(r, 30, True, NAVY)

    # ---------- 본문 요소 ----------
    def h2(self, text, container=None, before=2):
        c = container or self.doc
        p = c.add_paragraph(style='Heading 2') if container is None else c.add_paragraph()
        para_fmt(p, before, 4, 1.0, keep=True)
        r = p.add_run(text); set_run(r, 15, True, ORANGE)
        return p

    def label(self, text, container=None, before=4, color=NAVY, size=12):
        c = container or self.doc
        p = c.add_paragraph(); para_fmt(p, before, 2, 1.05, keep=True)
        r = p.add_run('■ '); set_run(r, size - 1, True, color)
        rich(p, text, size, color, True)
        return p

    def line(self, text, container=None, mark='', size=11, color='1F2933', indent=0, after=2, before=0):
        c = container or self.doc
        p = c.add_paragraph(); para_fmt(p, before, after, 1.12)
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        if mark:
            p.paragraph_format.left_indent = Cm(indent + 0.55)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            r = p.add_run(mark + ' '); set_run(r, size, True, ORANGE if mark in '①②③④⑤⑥⑦⑧⑨⑩' or mark[0].isdigit() else NAVY)
        rich(p, text, size, color)
        return p

    def bullets(self, items, container=None, mark='·', size=11, indent=0.2):
        for it in items:
            self.line(it, container, mark, size, indent=indent)

    def steps(self, items, container=None, size=11):
        nums = '①②③④⑤⑥⑦⑧⑨⑩'
        for i, it in enumerate(items):
            self.line(it, container, nums[i], size, after=3)

    def table(self, headers, rows, widths_cm, container=None, size=10, first_col_fill=True, center_cols=(), level_col=None):
        c = container or self.doc
        t = c.add_table(rows=1 + len(rows), cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        total = sum(widths_cm)
        avail = (BODY_W if container is None else None)
        widths = [Cm(w) for w in widths_cm]
        _fixed_layout(t, widths)
        for j, h in enumerate(headers):
            cell = t.rows[0].cells[j]
            _shade(cell, NAVY); _cell_margins(cell, 70, 70, 110, 110)
            _cell_borders(cell, top=(4, NAVY), bottom=(4, NAVY), left=(4, 'FFFFFF'), right=(4, 'FFFFFF'))
            p = cell.paragraphs[0]; para_fmt(p, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
            r = p.add_run(h); set_run(r, size, True, 'FFFFFF')
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _keep_row(t.rows[0])
        for i, row in enumerate(rows):
            _keep_row(t.rows[i + 1])
            for j, val in enumerate(row):
                cell = t.rows[i + 1].cells[j]
                _cell_margins(cell, 60, 60, 110, 110)
                _cell_borders(cell, top=(4, LINE), bottom=(4, LINE), left=(4, LINE), right=(4, LINE))
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if j == 0 and first_col_fill:
                    _shade(cell, NAVY_LIGHT)
                lines = val if isinstance(val, list) else [val]
                for k, ln in enumerate(lines):
                    p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                    para_fmt(p, 0, 0, 1.08, WD_ALIGN_PARAGRAPH.CENTER if j in center_cols else None)
                    if level_col is not None and j == level_col and ln in LEVEL:
                        lab, fg, bg = LEVEL[ln]
                        _shade(cell, bg)
                        r = p.add_run(lab); set_run(r, size, True, fg)
                    else:
                        rich(p, ln, size, '1F2933', j == 0 and first_col_fill)
        return t

    def callout(self, lines, kind='warn', container=None, size=11, width=None):
        c = container or self.doc
        bg, ln = {'warn': (WARN_BG, WARN_LINE), 'note': (NOTE_BG, LINE), 'ok': (OK_BG, 'B5D6C0')}[kind]
        t = c.add_table(rows=1, cols=1)
        _fixed_layout(t, [width or (BODY_W if container is None else Cm(14))])
        cell = t.rows[0].cells[0]
        _shade(cell, bg); _cell_margins(cell, 110, 110, 200, 200)
        _cell_borders(cell, top=(6, ln), bottom=(6, ln), left=(6, ln), right=(6, ln))
        for k, text in enumerate(lines):
            p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            para_fmt(p, 0, 1, 1.15)
            rich(p, text, size, '1F2933', k == 0 and kind == 'warn')
        sp = c.add_paragraph(); para_fmt(sp, 0, 0, 0.5)
        return t

    def cards(self, items, container=None, size=10.5, title_size=11.5, height=None):
        """items: (제목, [줄...]) 목록 → 가로 카드."""
        c = container or self.doc
        n = len(items)
        gap = Cm(0.35)
        w = (BODY_W - gap * (n - 1)) / n
        t = c.add_table(rows=1, cols=2 * n - 1)
        _no_table_borders(t)
        widths = []
        for i in range(n):
            widths.append(w)
            if i < n - 1:
                widths.append(gap)
        _fixed_layout(t, widths)
        _keep_row(t.rows[0])
        for i, (title, body) in enumerate(items):
            cell = t.rows[0].cells[2 * i]
            _cell_margins(cell, 120, 120, 170, 170)
            _cell_borders(cell, top=(24, NAVY), bottom=(4, LINE), left=(4, LINE), right=(4, LINE))
            p = cell.paragraphs[0]; para_fmt(p, 0, 4, 1.05)
            rich(p, title, title_size, NAVY, True)
            for b in body:
                bullet = not b.startswith('※') and not b.startswith('★')
                p = cell.add_paragraph(); para_fmt(p, 0, 2, 1.1)
                if bullet:
                    p.paragraph_format.left_indent = Cm(0.4)
                    p.paragraph_format.first_line_indent = Cm(-0.4)
                    r = p.add_run('· '); set_run(r, size, True, NAVY)
                rich(p, b, size, 'C0392B' if b.startswith('★') else '1F2933')
        if height:
            tr = t.rows[0]._tr.get_or_add_trPr()
            h = OxmlElement('w:trHeight'); h.set(qn('w:val'), str(int(height / 635))); h.set(qn('w:hRule'), 'atLeast')
            tr.append(h)
        sp = c.add_paragraph(); para_fmt(sp, 0, 0, 0.5)
        return t

    def two_col(self, left_w_cm, gap_cm=0.6):
        t = self.doc.add_table(rows=1, cols=3)
        _no_table_borders(t)
        lw = Cm(left_w_cm); g = Cm(gap_cm)
        _fixed_layout(t, [lw, g, BODY_W - lw - g])
        for cell in t.rows[0].cells:
            _cell_margins(cell, 0, 0, 0, 0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        L, R = t.rows[0].cells[0], t.rows[0].cells[2]
        # 빈 첫 문단 정리용 표시
        L._fresh = True; R._fresh = True
        return L, R

    def figure(self, path, width_cm, container=None, caption=None, frame=True):
        c = container or self.doc
        if frame:
            t = c.add_table(rows=1, cols=1)
            _fixed_layout(t, [Cm(width_cm)])
            cell = t.rows[0].cells[0]
            _shade(cell, 'E9ECF1'); _cell_margins(cell, 120, 120, 120, 120)
            _cell_borders(cell, top=(4, 'E9ECF1'), bottom=(4, 'E9ECF1'), left=(4, 'E9ECF1'), right=(4, 'E9ECF1'))
            p = cell.paragraphs[0]; para_fmt(p, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
            p.add_run().add_picture(path, width=Cm(width_cm - 0.45))
            if caption:
                p = cell.add_paragraph(); para_fmt(p, 4, 0, 1.05)
                rich(p, caption, 9.5, GREY)
            return
        else:
            p = c.add_paragraph(); para_fmt(p, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
            p.add_run().add_picture(path, width=Cm(width_cm))
        if caption:
            p = c.add_paragraph(); para_fmt(p, 3, 4, 1.05)
            rich(p, caption, 9.5, GREY)

    def badges(self, pairs, container=None, size=10):
        """블록 머리 정보: [(이름, 값), ...] 를 한 줄 표로."""
        c = container or self.doc
        n = len(pairs)
        t = c.add_table(rows=1, cols=2 * n)
        widths = []
        for k, v in pairs:
            widths += [Cm(2.6), (BODY_W - Cm(2.6) * n) / n]
        _fixed_layout(t, widths)
        for i, (k, v) in enumerate(pairs):
            kc, vc = t.rows[0].cells[2 * i], t.rows[0].cells[2 * i + 1]
            for cc in (kc, vc):
                _cell_margins(cc, 50, 50, 100, 100)
                _cell_borders(cc, top=(4, LINE), bottom=(4, LINE), left=(4, LINE), right=(4, LINE))
                cc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _shade(kc, NAVY_LIGHT)
            p = kc.paragraphs[0]; para_fmt(p, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
            r = p.add_run(k); set_run(r, size, True, NAVY)
            p = vc.paragraphs[0]; para_fmt(p, 0, 0, 1.0)
            if v in LEVEL:
                lab, fg, bg = LEVEL[v]
                _shade(vc, bg); r = p.add_run(lab); set_run(r, size, True, fg)
            else:
                rich(p, v, size)
        sp = c.add_paragraph(); para_fmt(sp, 0, 0, 0.6)

    def spacer(self, pt=4, container=None):
        c = container or self.doc
        p = c.add_paragraph(); para_fmt(p, 0, 0, 1.0)
        p.paragraph_format.space_after = Pt(pt)

    def save(self, path):
        # 셀의 빈 첫 문단 제거 (add_paragraph만 쓴 셀)
        for t in self.doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    ps = cell.paragraphs
                    if len(ps) > 1 and not ps[0].text and not ps[0]._p.xpath('.//w:drawing'):
                        ps[0]._p.getparent().remove(ps[0]._p)
        self.doc.save(path)
