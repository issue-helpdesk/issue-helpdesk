"""원래 스크린샷 자리가 없던 항목에도 그림과 그림 설명을 새로 넣는다 (사용자 가이드 docx·pptx).

- docx: 항목의 '조치 방법'/'신청 방법' 목록 바로 뒤에, 같은 문서의 기존 그림 표와 설명 문단을 본떠 넣는다.
- pptx: 항목의 마지막 장 뒤에 '화면 예시' 장을 추가한다.
이미 '그림 3.x.' 설명이 있는 항목은 건너뛴다. insert_figures.py를 먼저 실행해 둔 상태여야 한다.

실행: python dataset/guides/HW/tools/add_figures.py
"""
import copy
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation

HW = Path(__file__).resolve().parent.parent
NEXT_HEADINGS = ("해결 확인", "승인 절차")

CAPTIONS = {
    "IT-HW-001": {
        "3.3": "작업 관리자 [성능] 탭의 메모리 화면. 메모리 사용률을 확인하는 위치를 표시한다.",
        "3.4": "노트북 전원 연결 그림. 충전 표시등, 전원 버튼, 콘센트 위치를 표시한다.",
        "3.5": "파란 화면(블루스크린) 예시. 사진으로 찍어 둘 오류 문구 위치를 표시한다.",
        "3.7": "로그인 후 멈춘 화면 예시와 조치 순서.",
        "3.9": "듀얼 모니터 연결 그림. 다시 꽂을 케이블과 바꿔 볼 포트를 표시한다.",
        "3.10": "모니터 직접 연결 그림. 도킹스테이션 없이 노트북에 바로 연결하는 방법을 표시한다.",
        "3.11": "무선 마우스 수신기 위치 그림. 수신기를 꽂을 노트북 USB 포트를 표시한다.",
        "3.12": "OS 업데이트 '재시작 중' 화면 예시와 조치 순서.",
        "3.13": "소리 출력 장치 선택 화면. 출력 장치를 확인하는 위치를 표시한다.",
    },
    "IT-HW-002": {
        "3.1": "장비 신청 양식. '수리' 선택과 사고 경위·사진을 적는 위치를 표시한다.",
        "3.2": "장비 신청 양식. '수리' 선택과 증상을 적는 위치를 표시한다.",
        "3.4": "장비 신청 양식. '교체' 선택과 고사양 사유를 적는 위치, 승인 흐름을 표시한다.",
        "3.5": "장비 신청 양식. '신규 지급' 선택과 입사자 정보를 적는 위치를 표시한다.",
        "3.6": "장비 신청 양식. '비품' 선택과 주변기기 종류·증상을 적는 위치를 표시한다.",
        "3.7": "장비 신청 양식. '반납' 선택 위치를 표시한다.",
    },
    "IT-HW-003": {
        "3.2": "PC [설정] > [프린터 및 스캐너]의 복합기 화면. [제거] 버튼 위치를 표시한다.",
        "3.3": "복합기 오류 화면 예시와 전원 스위치. 적어 둘 오류 코드 위치를 표시한다.",
        "3.4": "복합기 '메일로 스캔' 화면과 메일함. 받는 메일 주소와 스팸(격리)함 위치를 표시한다.",
        "3.5": "복합기 공유폴더 저장 오류 화면 예시. '경로를 찾을 수 없음' 메시지를 표시한다.",
        "3.7": "복합기 요청 양식. 복합기 위치와 요청 내용을 적는 위치를 표시한다.",
    },
}


def _text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _style(el):
    ppr = el.find(qn("w:pPr"))
    st = ppr.find(qn("w:pStyle")) if ppr is not None else None
    return st.get(qn("w:val")) if st is not None else ""


def _set_caption(runs, no, body):
    runs[0].text = f"그림 {no}. "
    runs[1].text = body
    for r in runs[2:]:
        r.text = ""


def add_docx(path, doc_id):
    doc = Document(str(path))
    body = doc.element.body
    children = list(body.iterchildren())
    tpl_tbl = next(c for c in children if c.tag == qn("w:tbl") and c.findall(".//" + qn("w:drawing")))
    tpl_cap = children[children.index(tpl_tbl) + 1]
    width = Emu((int(tpl_tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))) - 260) * 635)
    existing = _text(body)
    n = 0
    for no, cap in CAPTIONS[doc_id].items():
        if f"그림 {no}. " in existing:
            continue
        children = list(body.iterchildren())
        start = next(i for i, c in enumerate(children)
                     if c.tag == qn("w:p") and _style(c) == "Heading2" and _text(c).startswith(no + " "))
        target = next(c for c in children[start + 1:]
                      if c.tag == qn("w:p") and _style(c) == "Heading3" and _text(c).strip() in NEXT_HEADINGS)
        tbl = copy.deepcopy(tpl_tbl)
        capp = copy.deepcopy(tpl_cap)
        target.addprevious(tbl)
        target.addprevious(capp)
        para = Table(tbl, doc._body).cell(0, 0).paragraphs[0]
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        full = f"그림 {no}. {cap}"
        pic = para.add_run().add_picture(str(HW / "images" / f"{doc_id}_fig{no}.png"), width=width)
        pic._inline.docPr.set("descr", full)
        _set_caption(Paragraph(capp, doc._body).runs, no, cap)
        n += 1
    if n:
        doc.save(str(path))
    return n


def _kicker(slide):
    sh = list(slide.shapes)[0]
    return sh.text_frame.text if sh.has_text_frame else ""


def add_pptx(path, doc_id):
    prs = Presentation(str(path))
    slides = list(prs.slides)
    tpl = next(s for s in slides if _kicker(s).endswith("화면 예시"))
    tpl_shapes = list(tpl.shapes)
    tpl_pic = next(s for s in tpl_shapes if s.shape_type == 13)
    tpl_cap = next(s for s in tpl_shapes if s.has_text_frame and s.text_frame.text.startswith("그림"))
    have = " ".join(s.text_frame.text for sl in slides for s in sl.shapes if s.has_text_frame)
    n = 0
    for no, cap in CAPTIONS[doc_id].items():
        if f"그림 {no}. " in have or f"그림 {no}-" in have:
            continue
        slides = list(prs.slides)
        idx = max(i for i, s in enumerate(slides) if f"· {no} " in _kicker(s))
        src = list(slides[idx].shapes)
        new = prs.slides.add_slide(slides[idx].slide_layout)
        for ph in list(new.placeholders):
            ph._element.getparent().remove(ph._element)
        for s in (src[0], src[1]):
            new.shapes._spTree.append(copy.deepcopy(s._element))
        runs = new.shapes[0].text_frame.paragraphs[0].runs
        runs[-1].text = re.sub(r"(\d+\.\d+) .*$", r"\1 화면 예시", runs[-1].text)
        full = f"그림 {no}. {cap}"
        pic = new.shapes.add_picture(str(HW / "images" / f"{doc_id}_fig{no}.png"),
                                     tpl_pic.left, tpl_pic.top, tpl_pic.width, tpl_pic.height)
        pic._element.nvPicPr.cNvPr.set("descr", full)
        new.shapes._spTree.append(copy.deepcopy(tpl_cap._element))
        _set_caption([r for p in new.shapes[-1].text_frame.paragraphs for r in p.runs], no, cap)
        new.shapes._spTree.append(copy.deepcopy(src[-1]._element))
        ids = prs.slides._sldIdLst
        el = ids[-1]
        ids.remove(el)
        ids.insert(idx + 1, el)
        n += 1
    if n:
        prs.save(str(path))
    return n


if __name__ == "__main__":
    for docx in sorted((HW / "guide").glob("*.docx")):
        doc_id = docx.name.split("_")[0]
        print(doc_id, "docx", add_docx(docx, doc_id), "pptx", add_pptx(docx.with_suffix(".pptx"), doc_id))
