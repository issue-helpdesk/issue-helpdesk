"""HW 사용자 가이드(docx·pptx)의 '[스크린샷 첨부 위치]' 자리에 images/ 의 그림을 넣는다.

- docx: 자리 표 안의 문구를 그림으로 바꾼다.
- pptx: 자리 상자에 그림을 넣는다. 자리가 작은 A·B 템플릿은 바로 다음 장에 같은 그림을 크게 넣은
  '화면 예시' 장을 추가한다.
그림 설명(캡션)은 고치지 않는다. 이미 그림이 들어간 파일은 자리 문구가 없으므로 건너뛴다.

실행: python dataset/guides/HW/tools/insert_figures.py
"""
import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Emu
from pptx import Presentation

HW = Path(__file__).resolve().parent.parent
MARKS = ("[스크린샷 첨부 위치]", "[가상 화면 그림 위치]")
CAPTION = re.compile(r"^그림 (\d+\.\d+)(?:-\d+)?\.")


def figure_for(doc_id, caption):
    m = CAPTION.match(caption.strip())
    path = HW / "images" / f"{doc_id}_fig{m.group(1)}.png"
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def fill_docx(path, doc_id):
    doc = Document(str(path))
    body = list(doc.element.body.iterchildren())
    n = 0
    for table in doc.tables:
        cell = table.cell(0, 0)
        if cell.text.strip() not in MARKS:
            continue
        nxt = body[body.index(table._tbl) + 1]
        caption = "".join(t.text or "" for t in nxt.iter(qn("w:t")))
        para = cell.paragraphs[0]
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        width = Emu((int(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))) - 260) * 635)
        pic = para.add_run().add_picture(str(figure_for(doc_id, caption)), width=width)
        pic._inline.docPr.set("descr", caption)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tr_pr = table.rows[0]._tr.trPr  # 빈 자리용 최소 높이를 없애 그림 높이에 맞춘다
        if tr_pr is not None:
            for hgt in tr_pr.findall(qn("w:trHeight")):
                tr_pr.remove(hgt)
        n += 1
    if n:
        doc.save(str(path))
    return n


def _fit(pic_w, pic_h, box_w, box_h):
    s = min(box_w / pic_w, box_h / pic_h)
    return int(pic_w * s), int(pic_h * s)


def fill_pptx(path, doc_id):
    prs = Presentation(str(path))
    targets = []
    for idx, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() in MARKS:
                targets.append((idx, slide, sh))
    layout = None
    added = 0
    for idx, slide, mark in targets:
        shapes = list(slide.shapes)
        cap = shapes[shapes.index(mark) + 1]
        fig = figure_for(doc_id, cap.text_frame.text)
        # 자리 상자 안의 작은 그림
        pad = 45720
        w, h = _fit(2400, 1040, mark.width - 2 * pad, mark.height - 2 * pad)
        small = slide.shapes.add_picture(str(fig), mark.left + (mark.width - w) // 2,
                                         mark.top + (mark.height - h) // 2, w, h)
        small._element.nvPicPr.cNvPr.set("descr", cap.text_frame.text)
        mark._element.getparent().remove(mark._element)
        if mark.height > 2000000:  # D 템플릿은 자리가 커서 그 안의 그림으로 충분하다
            continue

        # 다음 장: 큰 그림
        layout = slide.slide_layout
        new = prs.slides.add_slide(layout)
        for ph in list(new.placeholders):
            ph._element.getparent().remove(ph._element)
        kicker, title, number = shapes[0], shapes[1], shapes[-1]
        for src in (kicker, title):
            new.shapes._spTree.append(copy.deepcopy(src._element))
        k = new.shapes[0].text_frame.paragraphs[0].runs
        k[-1].text = re.sub(r"(\d+\.\d+) .*$", r"\1 화면 예시", k[-1].text)
        left, top, width = 457200, 987552, 8229600
        big = new.shapes.add_picture(str(fig), left, top, width, int(width * 1040 / 2400))
        big._element.nvPicPr.cNvPr.set("descr", cap.text_frame.text)
        c = copy.deepcopy(cap._element)
        new.shapes._spTree.append(c)
        moved = new.shapes[-1]
        moved.left, moved.top, moved.width, moved.height = left, 4590000, width, 201168
        new.shapes._spTree.append(copy.deepcopy(number._element))
        # 원래 장 바로 뒤로 옮긴다
        ids = prs.slides._sldIdLst
        el = ids[-1]
        ids.remove(el)
        ids.insert(idx + 1 + added, el)
        added += 1
    if targets:
        prs.save(str(path))
    return len(targets)


if __name__ == "__main__":
    for docx in sorted((HW / "guide").glob("*.docx")):
        doc_id = docx.name.split("_")[0]
        print(doc_id, "docx", fill_docx(docx, doc_id), "pptx", fill_pptx(docx.with_suffix(".pptx"), doc_id))
