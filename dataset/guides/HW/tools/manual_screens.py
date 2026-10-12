"""담당자 매뉴얼(-M)의 항목마다 '화면 예시' 그림을 만들어 docx·pptx에 넣는다.

왼쪽은 그 항목의 '처리 절차'에서 담당자가 쓰는 도구 화면(가상)이고, 오른쪽은 '처리 절차'를 docx에서 그대로
읽어 온 단계 카드다. 도구 이름은 매뉴얼 1장 '관련 시스템'의 가상 명칭을 쓴다. 값은 예시 자리표시만 넣는다.

- 그림 번호: 처리 흐름도는 '그림 3.x-1', 화면 예시는 '그림 3.x-2' (흐름도 설명의 번호를 이 스크립트가 바꾼다).
- docx: '처리 절차' 목록 바로 뒤. pptx: 처리 흐름도 장 바로 뒤에 '화면 예시' 장.
manual_flow.py를 먼저 실행해 둔 상태여야 한다. 이미 넣은 항목은 건너뛴다.

실행: python dataset/guides/HW/tools/manual_screens.py
"""
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import make_figures as mf  # noqa: E402
from make_figures import (BLUE, INK, LINE, NAVY, PANEL, RED, SUB, YELLOW, Fig, cable, dock, font, laptop,  # noqa: E402
                          monitor, ok_bad, stage, steps_card)
from manual_flow import _kicker, _style, _text, read_blocks  # noqa: E402

HW = Path(__file__).resolve().parent.parent
Y0, PH, LX, LW, RX, RW = 30, 920, 20, 1440, 1500, 880


# ---------------------------------------------------------------- 화면 부품
def body(f, x, y, w, spec):
    """도구 화면의 본문. spec: title, fields[(이름, 값, 딱지)], table(머리, 폭, 행, 표시 행, 딱지), buttons[(이름, 강조, 딱지)]."""
    f.text(x, y + 20, spec["title"], 32, bold=True, anchor="lm")
    f.hline(x, y + 56, w, fill="#9A9A9A")
    cy = y + 86
    for name, value, note in spec.get("fields", []):
        f.text(x, cy + 30, name, 24, bold=True, anchor="lm")
        fw = min(w - 250, 620)
        f.rect(x + 250, cy, fw, 60, fill="white", outline="#9A9A9A", r=6)
        f.text(x + 266, cy + 30, value, 24, SUB if value.startswith("(") else INK, anchor="lm")
        if note:
            f.mark((x + 250, cy, fw, 60), pad=6)
            f.label(x + 250, cy + 84, note)
            cy += 54
        cy += 82
    if "table" in spec:
        heads, widths, rows, hit, note = spec["table"]
        xs = [x + int(sum(widths[:i]) * w) for i in range(len(widths))]
        f.rect(x, cy, w, 54, fill="#EEF2F7")
        for hx, hd in zip(xs, heads):
            f.text(hx + 16, cy + 27, hd, 23, SUB, True, anchor="lm")
        cy += 54
        for i, row in enumerate(rows):
            if i == hit:
                f.rect(x, cy, w, 64, fill="#E8F1FB")
                f.mark((x, cy, w, 64), pad=2)
            for hx, v in zip(xs, row):
                bad = v.endswith("!")
                f.text(hx + 16, cy + 32, v.rstrip("!"), 24, "#A80000" if bad else INK, bad, anchor="lm")
            f.hline(x, cy + 64, w)
            cy += 64
        if note:
            f.label(x, cy + 22, note)
            cy += 70
        cy += 20
    bx = x
    for name, primary, note in spec.get("buttons", []):
        bw = 70 + len(name) * 26
        b = f.button(bx, cy + 10, bw, 62, name, primary=primary, size=25)
        if note:
            f.mark(b)
            f.label(bx + bw + 30, cy + 41, note, anchor="lm")
        bx += bw + 24
    return cy


def console(f, app, nav, sel, spec):
    """가상 관리 도구 창."""
    f.window(LX, Y0, LW, PH, app, bg="white")
    f.rect(LX + 2, Y0 + 60, 300, PH - 64, fill=NAVY)
    for i, n in enumerate(nav):
        ny = Y0 + 90 + i * 70
        if n == sel:
            f.rect(LX + 14, ny, 276, 56, fill="#3A5796", r=6)
        f.text(LX + 40, ny + 28, n, 24, "white", n == sel, anchor="lm")
    body(f, LX + 340, Y0 + 100, LW - 380, spec)


def remote(f, spec, inner_title):
    """원격 지원 도구 창 안에 사용자 PC 화면을 보여 준다."""
    f.window(LX, Y0, LW, PH, "원격 지원 도구", bg="#DDE3EA")
    f.rect(LX + 2, Y0 + 60, LW - 4, 60, fill="#107C10")
    f.text(LX + 30, Y0 + 90, "●  연결됨  ·  사용자 PC (EMP-0000)  ·  사용자 동의 완료", 24, "white", True, anchor="lm")
    ix, iy, iw, ih = LX + 60, Y0 + 150, LW - 120, PH - 190
    f.window(ix, iy, iw, ih, inner_title, bg="white")
    body(f, ix + 40, iy + 90, iw - 80, spec)


def safe_mode(f):
    f.rect(LX + 6, Y0 + 8, LW, PH, fill="#DADADA", r=14)
    f.rect(LX, Y0, LW, PH, fill="#101418", outline="#000000", width=3, r=14)
    for px, py, a in ((LX + 30, Y0 + 34, "lm"), (LX + LW - 30, Y0 + 34, "rm"), (LX + 30, Y0 + PH - 34, "lm"),
                      (LX + LW - 30, Y0 + PH - 34, "rm")):
        f.text(px, py, "안전 모드", 26, "white", True, anchor=a)


def wrap_lines(f, text, width=640, size=25):
    lines, cur = [], ""
    fnt = font(size)
    for word in text.split(" "):
        trial = f"{cur} {word}" if cur else word
        if cur and f.d.textlength(trial, font=fnt) > width:
            lines.append(cur)
            cur = word
        else:
            cur = trial
    return lines + [cur]


def procedure(f, block, hl):
    steps_card(f, RX, Y0, RW, PH, "처리 절차", [
        (i + 1, wrap_lines(f, s), "hl" if (i + 1) in hl else "") for i, s in enumerate(block["처리 절차"])])


# ---------------------------------------------------------------- 그림 정의
INV = ["자산 목록", "재고 현황", "반납·회수"]
MFP = ["복합기 목록", "오류 내역", "스캔·메일 설정", "소모품 상태"]
ASSET_FIELDS = [("사용자", "EMP-0000"), ("장비", "노트북 (모델명)"), ("지급일", "YYYY-MM-DD"),
                ("보증 만료일", "YYYY-MM-DD"), ("상태", "사용 중")]


def asset(note_on, note, buttons=()):
    return {"title": "자산 상세", "fields": [(n, v, note if n == note_on else None) for n, v in ASSET_FIELDS],
            "buttons": list(buttons)}


def stock(hit, note):
    return {"title": "재고 현황", "table": (["구분", "사양", "재고"], [0.34, 0.4, 0.26], [
        ["표준 노트북", "Windows", "(재고 수)"], ["개발용 노트북", "macOS / Windows", "(재고 수)"],
        ["고사양 노트북", "분석·AI 업무용", "(재고 수)"], ["대여 노트북", "수리 기간 임시 사용", "(재고 수)"],
        ["주변기기", "마우스 · 키보드", "(재고 수)"]], hit, note)}


def firmware(f, notes):
    f.window(LX, Y0, LW, PH, "도킹스테이션 펌웨어 업데이트 도구", bg=PANEL)
    dock(f, LX + 120, Y0 + 180, 360, 90, label="")
    body(f, LX + 120, Y0 + 340, LW - 240, {"title": "도킹스테이션 정보", "fields": [
        ("모델", "(도킹스테이션 모델)", None), ("현재 펌웨어", "(구버전)", notes.get("ver")),
        ("최신 펌웨어", "(최신 버전)", None)], "buttons": [("펌웨어 업데이트", True, notes.get("btn"))]})


def device_manager(f, x, y, w, h, group, device, button, note_dev, note_btn, title="장치 관리자"):
    f.window(x, y, w, h, title, bg="white")
    tx, ty = x + 40, y + 100
    for i, n in enumerate(["컴퓨터", "네트워크 어댑터", group, "키보드", "프로세서"]):
        f.text(tx, ty + i * 62, ("∨  " if n == group else ">  ") + n, 25, anchor="lm")
        if n == group:
            ty += 62
            f.rect(tx + 50, ty + i * 62 - 28, 520, 56, fill="#E8F1FB")
            f.d.polygon([(tx + 66, ty + i * 62 + 16), (tx + 94, ty + i * 62 + 16), (tx + 80, ty + i * 62 - 14)], fill="#F2A900")
            f.text(tx + 110, ty + i * 62, device, 25, anchor="lm")
            f.mark((tx + 50, ty + i * 62 - 28, 520, 56), pad=2)
            f.label(tx + 600, ty + i * 62, note_dev, anchor="lm")
    dx, dy, dw, dh = x + w - 760, y + h - 400, 700, 340
    f.window(dx, dy, dw, dh, f"{device} 속성", bg=PANEL)
    for i, n in enumerate(["일반", "드라이버", "자세히"]):
        f.rect(dx + 24 + i * 150, dy + 76, 146, 48, fill="white" if i == 1 else "#E6E6E6", outline=LINE)
        f.text(dx + 97 + i * 150, dy + 100, n, 23, anchor="mm")
    b = f.button(dx + 40, dy + 170, 300, 62, button)
    f.mark(b)
    f.label(dx + 40, dy + 268, note_btn)


def fig_adapter(f):
    stage(f, LX, Y0, LW, PH)
    lx, ly = LX + 110, Y0 + 230
    h = laptop(f, lx, ly)
    px, py = lx + 560, ly + h + 40
    f.d.ellipse([px + 6, py - 26, px + 26, py - 6], fill="#8A8A8A", outline="#555555", width=2)
    cable(f, [(px + 40, py + 10), (px + 190, py + 60), (px + 300, py - 40)])
    f.rect(px + 290, py - 90, 170, 80, fill="#3A4652", outline="#1B222C", width=3, r=10)
    f.text(px + 375, py - 50, "어댑터", 23, "white", anchor="mm")
    f.mark((px + 290, py - 90, 170, 80), pad=8)
    f.mark((px - 4, py - 36, 40, 40), pad=6)
    f.label(lx + 60, ly - 90, "1. 어댑터와 배터리 상태를 점검합니다 (충전 표시등 꺼짐)")
    f.label(px + 180, py + 60, "2. 불량 어댑터를 교체합니다")
    f.rect(px + 520, py - 110, 190, 120, fill="white", outline="#107C10", width=4, r=10)
    f.text(px + 615, py - 50, "교체용 어댑터", 23, "#107C10", True, anchor="mm")


def fig_cables(f):
    stage(f, LX, Y0, LW, PH)
    m1x, m2x, my, mw, mh = LX + 150, LX + 770, Y0 + 110, 520, 310
    monitor(f, m1x, my, mw, mh, "#0B5CAD", "모니터 1  (인식됨)")
    monitor(f, m2x, my, mw, mh, "#111111", "신호 없음", "#9AA3AE")
    dx, dy = LX + 590, Y0 + 700
    dock(f, dx, dy, label="노트북 · 도킹스테이션")
    cable(f, [(dx + 60, dy), (dx + 40, dy - 110), (m1x + mw / 2, my + mh + 62)])
    cable(f, [(dx + 200, dy), (dx + 230, dy - 110), (m2x + mw / 2, my + mh + 62)], color=RED, width=10)
    f.label(LX + 60, my + mh + 130, "1. 케이블 재연결과 [감지]를 안내합니다")
    f.rect(dx + 380, dy - 30, 330, 110, fill="white", outline="#107C10", width=4, r=10)
    f.text(dx + 545, dy + 25, "교체용 케이블", 24, "#107C10", True, anchor="mm")
    f.label(dx + 380, dy + 110, "2. 불량 케이블을 교체 지급합니다")


def fig_receiver(f):
    stage(f, LX, Y0, LW, PH)
    f.text(LX + 50, Y0 + 60, "끊기는 경우  (수신기가 도킹스테이션·모니터 포트에 있음)", 28, SUB, True, anchor="lm")
    ax, ay = LX + 150, Y0 + 230
    dock(f, ax, ay, 300, 80)
    mf._receiver(f, ax + 120, ay - 80)
    ok_bad(f, ax + 340, ay, False)
    f.hline(LX + 40, Y0 + 470, LW - 80, fill="#B5C2D3")
    bx, by = LX + 150, Y0 + 590
    h = laptop(f, bx, by, 300, "#0B5CAD")
    mf._receiver(f, bx + 372, by + h - 18)
    f.mark((bx + 360, by + h - 26, 58, 96), pad=6)
    ok_bad(f, bx + 460, by + h - 10, True)
    f.label(LX + 50, Y0 + 515, "1. 수신기를 노트북 본체 USB 포트로 옮기도록 안내합니다")
    f.label(bx + 600, by + 120, "2. 사용 중 끊김 여부를 확인합니다", anchor="lm")


def draw_screen(doc_id, block, out):
    no = block["no"]
    f = Fig()
    key = (doc_id, no)
    hl, cap = [1], ""
    if key == ("IT-HW-001-M", "3.1"):
        remote(f, {"title": "로컬 디스크 (C:) 폴더 크기", "table": (["폴더", "크기", "비고"], [0.5, 0.2, 0.3], [
            ["보안 에이전트 로그", "(큰 용량)!", "보관 기간 미설정"], ["임시 파일", "(용량)", ""],
            ["사용자 다운로드", "(용량)", ""], ["Windows", "(용량)", ""]], 0, "1. 원격 지원으로 C드라이브에서 큰 폴더를 확인"),
            "buttons": [("로그·임시 파일 정리", True, "2. 정리하고 로그 보관 기간 설정")]}, "파일 탐색기 - 폴더 크기")
        hl, cap = [1, 2], "원격 지원 도구 화면 예시. C드라이브에서 큰 폴더를 확인하고 정리하는 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.2"):
        remote(f, {"title": "시작 앱", "table": (["이름", "상태", "시작 시 영향"], [0.5, 0.22, 0.28], [
            ["사내 메신저", "사용", "높음"], ["음악 플레이어", "사용", "중간"], ["클라우드 동기화", "사용", "높음"],
            ["사내 보안 에이전트", "사용", "높음"]], 1, "1. 원격 지원으로 시작 프로그램과 디스크 사용률 확인"),
            "buttons": [("사용 안 함", False, "2. 불필요한 시작 프로그램 정리")]}, "작업 관리자")
        hl, cap = [1, 2], "원격 지원 도구 화면 예시. 사용자 PC의 시작 프로그램을 확인하고 정리하는 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.3"):
        remote(f, {"title": "성능", "table": (["항목", "사용률", "상태"], [0.4, 0.3, 0.3], [
            ["CPU", "34%", "정상"], ["메모리", "94%!", "부족!"], ["디스크 (SSD)", "41%", "읽기·쓰기 속도 점검"]], 1,
            "1. 원격 지원으로 메모리·디스크 상태를 점검")}, "작업 관리자")
        cap = "원격 지원 도구 화면 예시. 사용자 PC의 메모리·디스크 상태를 점검하는 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.4"):
        fig_adapter(f)
        hl, cap = [1, 2], "전원 어댑터 점검 그림. 점검할 어댑터와 충전 표시등, 교체 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.5"):
        safe_mode(f)
        device_manager(f, LX + 70, Y0 + 90, LW - 140, PH - 180, "디스플레이 어댑터", "그래픽 어댑터", "드라이버 롤백",
                       "문제가 있는 드라이버", "2. 그래픽 드라이버를 이전 버전으로 되돌림")
        f.label(LX + 240, Y0 + 34, "1. 안전 모드로 진입", anchor="lm")
        hl, cap = [1, 2], "안전 모드의 장치 관리자 화면 예시. [드라이버 롤백] 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.6"):
        console(f, "단말 관리 콘솔", ["단말 목록", "복구 키 조회", "정책"], "복구 키 조회", {
            "title": "디스크 암호화 복구 키 조회", "fields": [
                ("요청자", "EMP-0000  ·  본인 확인 완료", "1. 본인 확인을 한다"),
                ("복구 키 ID", "7C2A91E4-3B5D-4F08-A6C1-9E2D4B7F1A30", "2. 복구 키 ID로 복구 키를 조회"),
                ("복구 키", "●●●●●● - ●●●●●● - ●●●●●● - …", None)],
            "buttons": [("조회", True, None), ("복구 키 표시", False, "3. 복구 키를 전달하고 로그인까지 확인")]})
        hl, cap = [1, 2, 3], "단말 관리 콘솔 화면 예시. 복구 키 ID로 복구 키를 조회하는 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.7"):
        safe_mode(f)
        x, y, w, h = LX + 240, Y0 + 230, LW - 480, 440
        f.window(x, y, w, h, "시스템 파일 복구", bg=PANEL)
        f.text(x + 50, y + 130, "시스템 파일을 검사하고 손상된 파일을 복구하는 중입니다.", 25, anchor="lm")
        f.rect(x + 50, y + 190, w - 100, 40, fill="#E6E6E6", outline="#BDBDBD")
        f.rect(x + 50, y + 190, int((w - 100) * 0.6), 40, fill=BLUE)
        f.text(x + 50, y + 270, "진행 중…", 24, SUB, anchor="lm")
        f.mark((x + 50, y + 190, w - 100, 40), pad=8)
        f.label(x + 50, y + 330, "2. 시스템 파일을 복구한다")
        f.label(LX + 240, Y0 + 34, "1. 안전 모드로 진입", anchor="lm")
        f.label(LX + 240, Y0 + 760, "3. 복구가 안 되면 데이터 백업 후 OS 재설치")
        hl, cap = [1, 2, 3], "안전 모드의 시스템 파일 복구 화면 예시. 복구 진행 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.8"):
        firmware(f, {"ver": "1. 도킹스테이션 펌웨어 버전을 확인", "btn": "2. 펌웨어 업데이트 (안내 또는 원격)"})
        hl, cap = [1, 2], "도킹스테이션 펌웨어 업데이트 도구 화면 예시. 펌웨어 버전과 업데이트 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.9"):
        fig_cables(f)
        hl, cap = [1, 2], "듀얼 모니터 연결 그림. 다시 꽂도록 안내할 케이블과 교체 지급할 케이블을 표시한다."
    elif key == ("IT-HW-001-M", "3.10"):
        firmware(f, {"btn": "1. 도킹스테이션 펌웨어를 업데이트"})
        monitor(f, LX + 980, Y0 + 150, 320, 190, "#3B3B3B")
        f.label(LX + 940, Y0 + 104, "2. 패널 불량이면 모니터를 교체")
        hl, cap = [1, 2], "도킹스테이션 펌웨어 업데이트 도구 화면 예시. 펌웨어 업데이트 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.11"):
        fig_receiver(f)
        hl, cap = [1, 2], "무선 마우스 수신기 위치 그림. 사용자에게 안내할 수신기 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.12"):
        remote(f, {"title": "저장 공간과 업데이트", "fields": [
            ("로컬 디스크 (C:)", "여유 공간 부족", "1. 원격 또는 방문으로 디스크 공간을 확보"),
            ("OS 업데이트", "설치 실패 · 다시 시도 필요", None)],
            "buttons": [("업데이트 다시 시도", True, "2. 업데이트를 다시 진행")]}, "설정 - 시스템")
        hl, cap = [1, 2], "원격 지원 도구 화면 예시. 디스크 공간을 확보하고 업데이트를 다시 진행하는 위치를 표시한다."
    elif key == ("IT-HW-001-M", "3.13"):
        f.window(LX, Y0, LW, PH, "원격 지원 도구", bg="#DDE3EA")
        f.rect(LX + 2, Y0 + 60, LW - 4, 60, fill="#107C10")
        f.text(LX + 30, Y0 + 90, "●  연결됨  ·  사용자 PC (EMP-0000)  ·  사용자 동의 완료", 24, "white", True, anchor="lm")
        device_manager(f, LX + 60, Y0 + 150, LW - 120, PH - 190, "사운드, 비디오 및 게임 컨트롤러", "오디오 장치",
                       "드라이버 다시 설치", "1. 원격 지원으로 오디오 드라이버 확인", "2. 드라이버를 다시 설치")
        hl, cap = [1, 2], "원격 지원 도구의 장치 관리자 화면 예시. 오디오 드라이버를 확인하고 다시 설치하는 위치를 표시한다."
    elif key == ("IT-HW-002-M", "3.1"):
        console(f, "대여 장비 관리 대장", ["대여 현황", "대여 등록", "회수 등록"], "대여 현황", {
            "title": "대여 현황", "table": (["대여 장비", "상태", "사용자", "사유"], [0.28, 0.2, 0.22, 0.3], [
                ["대여 노트북 01", "대여 중", "EMP-0000", "침수·파손 수리"], ["대여 노트북 02", "대여 가능", "", ""],
                ["대여 노트북 03", "대여 가능", "", ""]], 0, "3. 수리 기간 동안 대여 노트북을 지급"),
            "buttons": [("회수 등록", False, "4. 수리 후 대여 장비를 회수")]})
        hl, cap = [3, 4], "대여 장비 관리 대장 화면 예시. 대여 노트북 지급과 회수를 기록하는 위치를 표시한다."
    elif key == ("IT-HW-002-M", "3.2"):
        console(f, "자산 관리 화면", INV, "자산 목록", asset("보증 만료일", "1. 자산 관리 화면에서 보증 기간을 확인",
                                                    [("보증 수리 접수", True, "2. 보증 수리를 접수")]))
        hl, cap = [1, 2], "자산 관리 화면 예시. 보증 기간을 확인하는 위치를 표시한다."
    elif key == ("IT-HW-002-M", "3.3"):
        console(f, "자산 관리 화면", INV, "자산 목록", asset("지급일", "1. 자산 관리 화면에서 지급일을 확인 (3년 경과 여부)"))
        cap = "자산 관리 화면 예시. 지급일을 확인하는 위치를 표시한다."
    elif key == ("IT-HW-002-M", "3.4"):
        console(f, "자산 관리 화면", INV, "재고 현황", stock(2, "2. 고사양 장비 재고를 확인"))
        hl, cap = [2], "자산 관리 화면의 재고 현황 예시. 고사양 장비 재고를 확인하는 위치를 표시한다."
    elif key == ("IT-HW-002-M", "3.5"):
        console(f, "자산 관리 화면", INV, "재고 현황", stock(0, "1. 재고를 확인"))
        cap = "자산 관리 화면의 재고 현황 예시. 신규 입사자용 장비 재고를 확인하는 위치를 표시한다."
    elif key == ("IT-HW-002-M", "3.6"):
        console(f, "자산 관리 화면", INV, "재고 현황", stock(4, "1. 재고를 확인"))
        cap = "자산 관리 화면의 재고 현황 예시. 주변기기 재고를 확인하는 위치를 표시한다."
    elif key == ("IT-HW-002-M", "3.7"):
        spec = asset("상태", "2. 장비를 회수해 자산 관리 화면에 반납 처리", [("반납 처리", True, None)])
        spec["fields"][-1] = ("상태", "반납", spec["fields"][-1][2])
        console(f, "자산 관리 화면", INV, "반납·회수", spec)
        hl, cap = [2], "자산 관리 화면 예시. 회수한 장비를 반납 처리하는 위치를 표시한다."
    elif key == ("IT-HW-003-M", "3.1"):
        remote(f, {"title": "서비스", "table": (["이름", "상태", "시작 유형"], [0.5, 0.25, 0.25], [
            ["인쇄 스풀러", "중지됨!", "자동"], ["원격 지원 에이전트", "실행 중", "자동"],
            ["사내 보안 에이전트", "실행 중", "자동"]], 0, "1. 원격 지원으로 인쇄 스풀러를 다시 시작"),
            "buttons": [("서비스 다시 시작", True, None), ("대기열 초기화", False, "1. 대기열을 초기화")]}, "서비스")
        cap = "원격 지원 도구 화면 예시. 인쇄 스풀러를 다시 시작하고 대기열을 초기화하는 위치를 표시한다."
    elif key == ("IT-HW-003-M", "3.2"):
        console(f, "복합기 관리 화면", MFP, "복합기 목록", {"title": "복합기 목록", "table": (
            ["복합기", "위치", "IP 주소", "상태"], [0.3, 0.22, 0.26, 0.22], [
                ["본사 3층 복합기", "본사 3층", "(현재 IP)!", "온라인"], ["본사 11층 복합기", "본사 11층", "(IP)", "온라인"],
                ["연구동 2층 복합기", "연구동 2층", "(IP)", "온라인"]], 0, "1. 복합기 관리 화면에서 현재 IP를 확인")})
        cap = "복합기 관리 화면 예시. 복합기의 현재 IP를 확인하는 위치를 표시한다."
    elif key == ("IT-HW-003-M", "3.3"):
        console(f, "복합기 관리 화면", MFP, "오류 내역", {"title": "오류 내역", "table": (
            ["복합기", "오류 코드", "발생 시각", "상태"], [0.3, 0.22, 0.26, 0.22], [
                ["본사 3층 복합기", "E-xx!", "(발생 시각)", "미해결!"], ["연구동 2층 복합기", "용지 걸림", "(발생 시각)", "해결"]],
            0, "1. 복합기 관리 화면에서 오류 내역을 확인"),
            "buttons": [("유지보수 업체 출동 요청", True, "2. 유지보수 업체에 출동을 요청")]})
        hl, cap = [1, 2], "복합기 관리 화면 예시. 오류 내역을 확인하고 업체 출동을 요청하는 위치를 표시한다."
    elif key == ("IT-HW-003-M", "3.4"):
        console(f, "복합기 관리 화면", MFP, "스캔·메일 설정", {"title": "메일 발송 설정  ·  본사 3층 복합기", "fields": [
            ("발송 상태", "메일 발송 오류 (인증 실패)", "1. 복합기 관리 화면에서 메일 발송 오류를 확인"),
            ("메일 서버(SMTP)", "(사내 메일 서버)", None), ("발송 계정", "(복합기 발송 계정)", None)],
            "buttons": [("계정 갱신", True, "2. 메일 발송(SMTP) 계정을 갱신"), ("테스트 발송", False, None)]})
        hl, cap = [1, 2], "복합기 관리 화면 예시. 메일 발송 오류를 확인하고 발송 계정을 갱신하는 위치를 표시한다."
    elif key == ("IT-HW-003-M", "3.5"):
        console(f, "복합기 관리 화면", MFP, "스캔·메일 설정", {"title": "스캔 저장 경로  ·  본사 3층 복합기", "fields": [
            ("등록된 경로", "(예전 공유폴더 경로)", None),
            ("새 경로", "(새 공유폴더 경로)", "2. 복합기 스캔 경로를 다시 설정")],
            "buttons": [("저장", True, None), ("테스트 스캔", False, "3. 테스트 스캔을 확인")]})
        hl, cap = [2, 3], "복합기 관리 화면 예시. 스캔 저장 경로를 다시 설정하는 위치를 표시한다."
    elif key == ("IT-HW-003-M", "3.6"):
        f.window(LX, Y0, LW, PH, "소프트웨어 센터")
        f.rect(LX + 2, Y0 + 60, 300, PH - 64, fill=NAVY)
        for i, n in enumerate(["응용 프로그램", "프린터 설치", "설치 상태"]):
            if i == 1:
                f.rect(LX + 14, Y0 + 90 + i * 70, 276, 56, fill="#3A5796", r=6)
            f.text(LX + 40, Y0 + 118 + i * 70, n, 24, "white", i == 1, anchor="lm")
        body(f, LX + 340, Y0 + 100, LW - 380, {"title": "프린터 설치", "table": (["복합기", "위치", ""], [0.4, 0.35, 0.25], [
            ["본사 3층 복합기", "본사 · 3층", "[설치]"], ["본사 11층 복합기", "본사 · 11층", "[설치]"],
            ["연구동 2층 복합기", "연구동 · 2층", "[설치]"]], 0, "1. 설치 방법을 안내 (사용자 가이드 IT-HW-003 3.6)"),
            "buttons": [("테스트 페이지 인쇄", False, "2. 테스트 인쇄를 확인")]})
        hl, cap = [1, 2], "소프트웨어 센터 [프린터 설치] 화면 예시. 사용자에게 안내할 설치 위치를 표시한다."
    elif key == ("IT-HW-003-M", "3.7"):
        console(f, "복합기 관리 화면", MFP, "소모품 상태", {"title": "소모품 상태", "table": (
            ["복합기", "토너", "급지 롤러", "용지 걸림"], [0.3, 0.22, 0.24, 0.24], [
                ["본사 3층 복합기", "교체 필요!", "마모!", "반복!"], ["본사 11층 복합기", "정상", "정상", "없음"],
                ["연구동 2층 복합기", "정상", "정상", "없음"]], 0, "1. 토너를 교체"),
            "buttons": [("롤러 교체 요청", True, "2. 급지 롤러 마모면 롤러 교체를 요청")]})
        hl, cap = [1, 2], "복합기 관리 화면 예시. 소모품 상태를 확인하고 롤러 교체를 요청하는 위치를 표시한다."
    else:
        raise KeyError(key)
    procedure(f, block, hl)
    mf.OUT = out.parent
    f.save(out.name)
    return cap


# ---------------------------------------------------------------- 문서에 넣기
def _renumber(runs_list, numbers):
    """처리 흐름도 설명 '그림 3.x. ' → '그림 3.x-1. '"""
    for runs in runs_list:
        if runs and re.fullmatch(r"그림 (\d+\.\d+)\. ", runs[0].text or "") and "처리 흐름도" in "".join(r.text for r in runs):
            runs[0].text = runs[0].text.replace(". ", "-1. ")


def fill_docx(path, doc_id):
    doc = Document(str(path))
    _renumber([Paragraph(p, doc._body).runs for p in doc.element.body.iter(qn("w:p"))], None)
    blocks = read_blocks(doc)
    caps = {b["no"]: draw_screen(doc_id, b, HW / "images" / f"{doc_id}_fig{b['no']}-2.png") for b in blocks}
    children = list(doc.element.body.iterchildren())
    tpl_tbl = next(c for c in children if c.tag == qn("w:tbl") and c.findall(".//" + qn("w:drawing")))
    tpl_cap = children[children.index(tpl_tbl) + 1]
    width = Emu((int(tpl_tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))) - 260) * 635)
    existing = _text(doc.element.body)
    n = 0
    for b in blocks:
        no = b["no"]
        if f"그림 {no}-2. " in existing:
            continue
        children = list(doc.element.body.iterchildren())
        start = children.index(b["info_tbl"])
        target = next(c for c in children[start:] if c.tag == qn("w:p") and _style(c) == "Heading3"
                      and _text(c).strip() == "처리 후 안내 문구")
        tbl, capp = copy.deepcopy(tpl_tbl), copy.deepcopy(tpl_cap)
        target.addprevious(tbl)
        target.addprevious(capp)
        para = Table(tbl, doc._body).cell(0, 0).paragraphs[0]
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        pic = para.add_run().add_picture(str(HW / "images" / f"{doc_id}_fig{no}-2.png"), width=width)
        pic._inline.docPr.set("descr", f"그림 {no}-2. {caps[no]}")
        runs = Paragraph(capp, doc._body).runs
        runs[0].text, runs[1].text = f"그림 {no}-2. ", caps[no]
        n += 1
    doc.save(str(path))
    return caps, n


def fill_pptx(path, doc_id, caps):
    prs = Presentation(str(path))
    _renumber([[r for p in s.text_frame.paragraphs for r in p.runs] for sl in prs.slides for s in sl.shapes
               if s.has_text_frame], None)
    have = " ".join(s.text_frame.text for sl in prs.slides for s in sl.shapes if s.has_text_frame)
    n = 0
    for no, cap in caps.items():
        if f"그림 {no}-2. " in have:
            continue
        slides = list(prs.slides)
        idx = max(i for i, s in enumerate(slides) if f"· {no} " in _kicker(s))   # 처리 흐름도 장
        src = list(slides[idx].shapes)
        new = prs.slides.add_slide(slides[idx].slide_layout)
        for ph in list(new.placeholders):
            ph._element.getparent().remove(ph._element)
        for s in (src[0], src[1]):
            new.shapes._spTree.append(copy.deepcopy(s._element))
        runs = new.shapes[0].text_frame.paragraphs[0].runs
        runs[-1].text = re.sub(r"(\d+\.\d+) .*$", r"\1 화면 예시", runs[-1].text)
        spic = next(s for s in src if s.shape_type == 13)
        scap = next(s for s in src if s.has_text_frame and s.text_frame.text.startswith("그림"))
        pic = new.shapes.add_picture(str(HW / "images" / f"{doc_id}_fig{no}-2.png"), spic.left, spic.top, spic.width,
                                     spic.height)
        pic._element.nvPicPr.cNvPr.set("descr", f"그림 {no}-2. {cap}")
        new.shapes._spTree.append(copy.deepcopy(scap._element))
        cr = [r for p in new.shapes[-1].text_frame.paragraphs for r in p.runs]
        cr[0].text, cr[1].text = f"그림 {no}-2. ", cap
        new.shapes._spTree.append(copy.deepcopy(src[-1]._element))
        ids = prs.slides._sldIdLst
        el = ids[-1]
        ids.remove(el)
        ids.insert(idx + 1, el)
        n += 1
    prs.save(str(path))
    return n


if __name__ == "__main__":
    for docx in sorted((HW / "manual").glob("*.docx")):
        doc_id = docx.name.split("_")[0]
        caps, a = fill_docx(docx, doc_id)
        print(doc_id, "항목", len(caps), "docx", a, "pptx", fill_pptx(docx.with_suffix(".pptx"), doc_id, caps))
