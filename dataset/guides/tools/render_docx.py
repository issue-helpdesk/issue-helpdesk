"""원천 JSON → DOCX (편집용, A4 세로). 색·표지·장 구성은 장표(PDF·PPTX)와 같게 맞춘다."""
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

import layout as LY
from model import COMPANY, COPYRIGHT, CATEGORIES

FACE = LY.FONT_FACE
CONTENT_W = 17.0   # cm


def rgb(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_font(run, size=None, bold=None, color=None, mono=False):
    name = LY.MONO if mono else FACE
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = rgb(color)


def add_runs(par, text, size=10.5, color=LY.INK, bold=False):
    for t, b, c in LY.runs(text, color, bold):
        r = par.add_run(t)
        set_font(r, size, b, c)
    return par


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcpr.append(shd)
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill.lstrip('#'))


def borders(tbl, color=LY.LINE, sz=4):
    tblpr = tbl._tbl.tblPr
    b = tblpr.find(qn('w:tblBorders'))
    if b is None:
        b = OxmlElement('w:tblBorders')
        tblpr.append(b)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = b.find(qn(f'w:{edge}'))
        if e is None:
            e = OxmlElement(f'w:{edge}')
            b.append(e)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:color'), color.lstrip('#'))


def cell_margins(tbl, top=60, side=100):
    tblpr = tbl._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for edge, v in (('top', top), ('left', side), ('bottom', top), ('right', side)):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tblpr.append(m)


def set_widths(tbl, widths_cm):
    tbl.autofit = False
    grid = tbl._tbl.tblGrid
    for j, w in enumerate(widths_cm):
        if j < len(grid.gridCol_lst):
            grid.gridCol_lst[j].set(qn('w:w'), str(int(w * 567)))
        for r in tbl.rows:
            if j < len(r.cells):
                r.cells[j].width = Cm(w)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader')
    e.set(qn('w:val'), 'true')
    trpr.append(e)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit')
    e.set(qn('w:val'), 'true')
    trpr.append(e)


def cell_text(cell, text, size=9.5, color=LY.INK, bold=False, align=None):
    cell.text = ''
    lines = str(text).split('\n')
    for n, ln in enumerate(lines):
        par = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.line_spacing = 1.15
        if align:
            par.alignment = align
        add_runs(par, ln, size, color, bold)


class W:
    def __init__(self, doc, img_dir):
        self.src, self.m, self.img_dir = doc, doc['meta'], img_dir
        self.d = Document()
        self.fig_sec, self.fig_no = '', 0
        self._setup()

    # ------------------------------------------------ 기본 설정
    def _setup(self):
        d = self.d
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
        sec.header_distance = sec.footer_distance = Cm(1)
        st = d.styles['Normal']
        st.font.name = FACE
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), FACE)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.25
        for lvl, size, color in ((1, 18, LY.INK), (2, 14, LY.NAVY), (3, 12, LY.SUB)):
            hs = d.styles[f'Heading {lvl}']
            hs.font.name = FACE
            hs.element.rPr.rFonts.set(qn('w:eastAsia'), FACE)
            hs.element.rPr.rFonts.set(qn('w:ascii'), FACE)
            hs.element.rPr.rFonts.set(qn('w:hAnsi'), FACE)
            hs.font.size = Pt(size)
            hs.font.bold = True
            hs.font.color.rgb = rgb(color)
            hs.paragraph_format.space_before = Pt(14 if lvl < 3 else 8)
            hs.paragraph_format.space_after = Pt(5)
            hs.paragraph_format.keep_with_next = True
        cp = d.core_properties
        cp.title = f"{self.m['doc_no']} {self.m['title']}"
        cp.author = self.m['owner']
        cp.last_modified_by = self.m['owner']
        cp.subject = f"{self.m['access']} · {self.m.get('audience', '')}"
        cp.keywords = ', '.join(self.m.get('keywords', []))[:250]
        cp.category = self.m['access']
        self._header_footer(sec)

    def _header_footer(self, sec):
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.text = ''
        add_runs(hp, f"{self.m['doc_no']}  {self.m['title']}", 8.5, LY.MUTED)
        r = hp.add_run(f"   |   {self.m['access']}")
        set_font(r, 8.5, True, LY.SUB if self.m['access'] != '사내 일반' else LY.NAVY)
        self._bottom_border(hp, LY.NAVY)
        fp = sec.footer.paragraphs[0]
        fp.text = ''
        add_runs(fp, f"{COPYRIGHT} · {self.m['doc_no']} · {self.m['access']}", 8, LY.MUTED)
        r = fp.add_run('   |   ')
        set_font(r, 8, False, LY.MUTED)
        self._field(fp, 'PAGE')
        r = fp.add_run(' / ')
        set_font(r, 8, False, LY.MUTED)
        self._field(fp, 'NUMPAGES')

    def _field(self, par, code):
        r = par.add_run()
        set_font(r, 8, False, LY.MUTED)
        for tag, attr in (('w:fldChar', 'begin'), ('w:instrText', None), ('w:fldChar', 'separate'), ('w:t', None), ('w:fldChar', 'end')):
            e = OxmlElement(tag)
            if tag == 'w:fldChar':
                e.set(qn('w:fldCharType'), attr)
            elif tag == 'w:instrText':
                e.set(qn('xml:space'), 'preserve')
                e.text = f' {code} '
            else:
                e.text = '1'
            r._r.append(e)

    def _bottom_border(self, par, color, sz=8):
        ppr = par._p.get_or_add_pPr()
        b = OxmlElement('w:pBdr')
        e = OxmlElement('w:bottom')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '4')
        e.set(qn('w:color'), color.lstrip('#'))
        b.append(e)
        ppr.append(b)

    def _top_bar(self, par, color, sz=24):
        ppr = par._p.get_or_add_pPr()
        b = OxmlElement('w:pBdr')
        e = OxmlElement('w:top')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '12')
        e.set(qn('w:color'), color.lstrip('#'))
        b.append(e)
        ppr.append(b)

    # ------------------------------------------------ 표지·문서 정보·목차
    def cover(self):
        m = self.m
        d = self.d
        for _ in range(6):
            d.add_paragraph()
        bar = d.add_paragraph()
        bar.paragraph_format.right_indent = Cm(14.5)
        self._top_bar(bar, LY.ORANGE)
        from model import series as _series, RESTRICTED_NOTE
        series = _series(m)
        add_runs(d.add_paragraph(), series, 12, LY.MUTED)
        t = d.add_paragraph()
        add_runs(t, m['title'], 28, LY.INK, True)
        t.paragraph_format.line_spacing = 1.1
        add_runs(d.add_paragraph(), f"{m['doc_no']}  |  {m['owner']}  |  {m['date'].replace('-', '.')}  |  Ver {m['version'].lstrip('v')}", 12, LY.NAVY, True)
        restricted = m['access'] != '사내 일반'
        tb = d.add_table(rows=1, cols=1)
        c = tb.cell(0, 0)
        shade(c, LY.SUB if restricted else LY.NAVY)
        cell_text(c, f"{m['access']} · {m.get('audience', '')}", 10, LY.WHITE, True)
        set_widths(tb, [6.5])
        if restricted:
            d.add_paragraph()
            add_runs(d.add_paragraph(), RESTRICTED_NOTE.get(m['doc_type'], RESTRICTED_NOTE['manual']), 9.5, LY.MUTED)
        self.heading_page('문서 정보', new_page=True)
        rows = LY.doc_info_rows(m)
        tb = d.add_table(rows=len(rows), cols=4)
        borders(tb)
        cell_margins(tb)
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = tb.cell(i, j)
                key = j in (0, 2)
                shade(c, LY.KEY if key else LY.CELL)
                cell_text(c, v, 9.5, LY.NAVY if key else LY.INK, key)
        set_widths(tb, [2.8, 5.7, 2.8, 5.7])
        if m.get('keywords'):
            d.add_paragraph()
            tb = d.add_table(rows=1, cols=2)
            borders(tb)
            cell_margins(tb)
            shade(tb.cell(0, 0), LY.KEY)
            cell_text(tb.cell(0, 0), '키워드', 9.5, LY.NAVY, True)
            cell_text(tb.cell(0, 1), ', '.join(m['keywords']), 9.5)
            set_widths(tb, [2.8, 14.2])
        d.add_paragraph()
        self.heading_page('Contents')
        self._toc()

    def heading_page(self, text, new_page=False):
        p_ = self.d.add_paragraph()
        p_.paragraph_format.page_break_before = new_page
        add_runs(p_, text, 18, LY.INK, True)
        self._bottom_border(p_, LY.NAVY, 12)
        p_.paragraph_format.space_after = Pt(10)

    def _toc(self):
        """Word 목차 필드. 장·항목 목록을 미리 넣어 두고, [필드 업데이트]를 하면 쪽 번호가 붙는다."""
        par = self.d.add_paragraph()
        r = par.add_run()
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), 'begin')
        r._r.append(fc)
        r = par.add_run()
        it = OxmlElement('w:instrText')
        it.set(qn('xml:space'), 'preserve')
        it.text = ' TOC \\o "1-2" \\h \\z \\u '
        r._r.append(it)
        r = par.add_run()
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), 'separate')
        r._r.append(fc)
        first = True
        for ch in self.src['chapters']:
            label = f"{ch['no']}. {ch['title']}" if ch['no'].isdigit() else f"{ch['no']} · {ch['title']}"
            p2 = par if first else self.d.add_paragraph()
            first = False
            add_runs(p2, label, 11, LY.NAVY, True)
            for s in ch['sections']:
                p3 = self.d.add_paragraph()
                p3.paragraph_format.left_indent = Cm(0.8)
                add_runs(p3, f"{s['no']} {s['title']}", 9.5, LY.MUTED)
        p4 = self.d.add_paragraph()
        add_runs(p4, '개정 이력', 11, LY.NAVY, True)
        r = p4.add_run()
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), 'end')
        r._r.append(fc)

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------ 본문
    def chapter(self, ch):
        label = f"{ch['no']}. {ch['title']}" if ch['no'].isdigit() else f"{ch['no']} · {ch['title']}"
        hp = self.d.add_heading(level=1)
        hp.text = ''
        hp.paragraph_format.page_break_before = True
        if ch['no'].isdigit():
            r = hp.add_run(f"{int(ch['no']):02d}  ")
            set_font(r, 18, True, LY.ORANGE)
            r = hp.add_run(ch['title'])
        else:
            r = hp.add_run(label)
        set_font(r, 18, True, LY.INK)
        self._bottom_border(hp, LY.NAVY, 12)
        self.fig_sec, self.fig_no = (ch['no'] if ch['no'].isdigit() else ''), 0
        for it in ch['items']:
            self.item(it)
        for s in ch['sections']:
            self.fig_sec, self.fig_no = s['no'], 0
            h2 = self.d.add_heading(level=2)
            h2.text = ''
            r = h2.add_run(f"{s['no']} {s['title']}")
            set_font(r, 14, True, LY.NAVY)
            if s.get('props'):
                self.props(s['props'])
            for it in s['items']:
                self.item(it)

    def props(self, rows):
        tb = self.d.add_table(rows=1, cols=len(rows) * 2)
        borders(tb)
        cell_margins(tb, 40, 80)
        ws = []
        for j, (k, v) in enumerate(rows):
            c1, c2 = tb.cell(0, 2 * j), tb.cell(0, 2 * j + 1)
            shade(c1, LY.KEY)
            cell_text(c1, k, 8.5, LY.NAVY, True)
            fill, color = LY.BADGE.get(v, (LY.CELL, LY.INK))
            shade(c2, fill)
            cell_text(c2, v, 8.5, color, True)
            ws += [len(k) * 0.32 + 0.4, len(v) * 0.3 + 0.5]
        tot = sum(ws)
        set_widths(tb, [w * CONTENT_W / tot for w in ws] if tot > CONTENT_W else ws)
        self.d.add_paragraph().paragraph_format.space_after = Pt(0)

    def item(self, it):
        t = it['t']
        d = self.d
        if t == 'h':
            h = d.add_heading(level=3)
            h.text = ''
            r = h.add_run(it['text'])
            set_font(r, 12, True, LY.SUB)
        elif t == 'p':
            add_runs(d.add_paragraph(), it['text'])
        elif t in ('ul', 'ol'):
            for n, x in enumerate(it['items'], 1):
                par = d.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.7)
                par.paragraph_format.first_line_indent = Cm(-0.5)
                par.paragraph_format.space_after = Pt(2)
                mark = (LY._circ(n) if t == 'ol' else '•') + ' '
                r = par.add_run(mark)
                set_font(r, 10.5, True, LY.ORANGE if t == 'ol' else LY.NAVY)
                add_runs(par, x)
        elif t == 'table':
            self.table(it['head'], it['rows'], it.get('widths'), it.get('badge_col'))
        elif t == 'kv':
            self.table(None, it['rows'], [1, 3.2], key_col=True)
        elif t == 'qa':
            self.table(['질문', '답변'], it['items'], [4, 8], key_col=True)
        elif t == 'cols':
            n = len(it['cols'])
            tb = d.add_table(rows=2, cols=n)
            borders(tb)
            cell_margins(tb)
            for j, c in enumerate(it['cols']):
                shade(tb.cell(0, j), LY.KEY)
                cell_text(tb.cell(0, j), c['title'], 10, LY.NAVY, True)
                cell_text(tb.cell(1, j), '\n'.join('· ' + x for x in c['items']), 9.5)
            set_widths(tb, [CONTENT_W / n] * n)
            d.add_paragraph().paragraph_format.space_after = Pt(0)
        elif t == 'note':
            star = it['kind'] == 'star'
            tb = d.add_table(rows=1, cols=1)
            c = tb.cell(0, 0)
            shade(c, LY.STAR_BG if star else LY.INFO_BG)
            borders(tb, LY.STAR_LINE if star else LY.INFO_BG)
            cell_margins(tb, 80, 140)
            cell_text(c, ('★ ' if star else '※ ') + it['title'], 10, LY.SUB if star else LY.MUTED, star)
            for ln in it.get('lines', []):
                par = c.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                add_runs(par, ln, 9.5)
            set_widths(tb, [CONTENT_W])
            d.add_paragraph().paragraph_format.space_after = Pt(0)
        elif t == 'quote':
            tb = d.add_table(rows=1, cols=1)
            c = tb.cell(0, 0)
            shade(c, LY.WHITE)
            borders(tb, LY.LINE)
            tcpr = c._tc.get_or_add_tcPr()
            b = OxmlElement('w:tcBorders')
            e = OxmlElement('w:left')
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), '24')
            e.set(qn('w:color'), LY.NAVY.lstrip('#'))
            b.append(e)
            tcpr.append(b)
            cell_margins(tb, 80, 160)
            cell_text(c, it['text'], 10)
            set_widths(tb, [CONTENT_W])
            d.add_paragraph().paragraph_format.space_after = Pt(0)
        elif t == 'code':
            tb = d.add_table(rows=1, cols=1)
            c = tb.cell(0, 0)
            shade(c, LY.CODE_BG)
            cell_margins(tb, 100, 160)
            c.text = ''
            for n, ln in enumerate(it['text'].split('\n')):
                par = c.paragraphs[0] if n == 0 else c.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                r = par.add_run(ln)
                set_font(r, 9, False, LY.CODE_FG, mono=True)
            set_widths(tb, [CONTENT_W])
            d.add_paragraph().paragraph_format.space_after = Pt(0)
        elif t == 'flow':
            n = len(it['steps'])
            tb = d.add_table(rows=3, cols=n)
            borders(tb)
            cell_margins(tb)
            for j, s in enumerate(it['steps']):
                cell_text(tb.cell(0, j), f'STEP {j + 1}', 8.5, LY.ORANGE, True, WD_ALIGN_PARAGRAPH.CENTER)
                shade(tb.cell(1, j), LY.NAVY)
                cell_text(tb.cell(1, j), s['title'], 10, LY.WHITE, True, WD_ALIGN_PARAGRAPH.CENTER)
                cell_text(tb.cell(2, j), s['desc'], 9, LY.INK, False, WD_ALIGN_PARAGRAPH.CENTER)
            set_widths(tb, [CONTENT_W / n] * n)
            d.add_paragraph().paragraph_format.space_after = Pt(0)
        elif t == 'fig':
            self.figure(it)

    def table(self, head, rows, widths=None, badge_col=None, key_col=False):
        ncol = len(head) if head else len(rows[0])
        widths = widths or LY._auto_widths(head, rows, ncol, 9.5)
        tot = sum(widths)
        wcm = [CONTENT_W * w / tot for w in widths]
        tb = self.d.add_table(rows=(1 if head else 0) + len(rows), cols=ncol)
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(tb)
        cell_margins(tb)
        i0 = 0
        if head:
            for j, h in enumerate(head):
                c = tb.cell(0, j)
                shade(c, LY.NAVY)
                cell_text(c, h, 9.5, LY.WHITE, True, WD_ALIGN_PARAGRAPH.CENTER)
            repeat_header(tb.rows[0])
            i0 = 1
        for i, r in enumerate(rows):
            r = list(r) + [''] * (ncol - len(r))
            for j in range(ncol):
                c = tb.cell(i0 + i, j)
                v = str(r[j])
                fill, color, bold, al = LY.CELL, LY.INK, False, None
                if key_col and j == 0:
                    fill, color, bold = LY.KEY, LY.NAVY, True
                if badge_col is not None and j == badge_col and v in LY.BADGE:
                    fill, color = LY.BADGE[v]
                    bold, al = True, WD_ALIGN_PARAGRAPH.CENTER
                shade(c, fill)
                cell_text(c, v, 9.5, color, bold, al)
            no_split(tb.rows[i0 + i])
        set_widths(tb, wcm)
        self.d.add_paragraph().paragraph_format.space_after = Pt(0)

    def figure(self, f):
        path = os.path.join(self.img_dir, f['src'])
        if not os.path.exists(path):
            return
        self.fig_no += 1
        iw, ih = Image.open(path).size
        w = min(15.5, 15.5 * (iw / ih) / 1.2) if ih > iw * 0.85 else 15.5
        w = max(w, 7)
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Cm(w))
        label = f"{self.fig_sec}-{self.fig_no}" if self.fig_sec else str(self.fig_no)
        cap = self.d.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(cap, f"그림 {label}. {f.get('caption', '')}".rstrip('. ').rstrip(), 9, LY.MUTED)
        side = ([f['step']] if f.get('step') else []) + list(f.get('points', []))
        for k, s in enumerate(side):
            par = self.d.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.5)
            par.paragraph_format.space_after = Pt(1)
            add_runs(par, s, 9.5, LY.INK, k == 0 and bool(f.get('step')))
        if f.get('note'):
            par = self.d.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.5)
            add_runs(par, f['note'], 9.5, LY.SUB)

    def history(self):
        h = self.d.add_heading(level=1)
        h.text = ''
        h.paragraph_format.page_break_before = True
        r = h.add_run('개정 이력')
        set_font(r, 18, True, LY.INK)
        self._bottom_border(h, LY.NAVY, 12)
        self.table(['버전', '개정일', '작성자', '변경 내용'], self.m['history'], [1, 1.4, 1.6, 7])

    def build(self, path):
        self.cover()
        for ch in self.src['chapters']:
            self.chapter(ch)
        self.history()
        self.d.save(path)


def render(doc, img_dir, path):
    W(doc, img_dir).build(path)
