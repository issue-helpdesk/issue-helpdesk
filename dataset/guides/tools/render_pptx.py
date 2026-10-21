"""layout 결과를 PPTX로 만든다. PDF와 같은 좌표·같은 줄바꿈을 쓰고, 표는 PowerPoint 표 개체로 만든다."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt, Emu

import layout as LY

EMU = 12700  # 1pt


def P(v):
    return Emu(int(round(v * EMU)))


def rgb(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(run, size, bold, color, mono=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bool(bold)
    f.color.rgb = rgb(color)
    name = LY.MONO if mono else LY.FONT_FACE
    f.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set('typeface', name)


def _fill_tf(tf, lines, size, align, leading, mono=False):
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    al = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
    for n, ln in enumerate(lines):
        para = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        para.alignment = al
        para.line_spacing = Pt(leading)
        para.space_before = para.space_after = Pt(0)
        if not ln:
            ln = [(' ', False, LY.INK)]
        for t, b, c in ln:
            r = para.add_run()
            r.text = t
            _set_font(r, size, b, c, mono)


def add_text(slide, el):
    h = max(len(el['lines']), 1) * el['leading']
    tb = slide.shapes.add_textbox(P(el['x']), P(el['y']), P(el['w']), P(h))
    _fill_tf(tb.text_frame, el['lines'], el['size'], el['align'], el['leading'], el.get('mono'))


def add_rect(slide, el):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if el.get('radius') else MSO_SHAPE.RECTANGLE,
                                 P(el['x']), P(el['y']), P(el['w']), P(el['h']))
    if el.get('radius'):
        shp.adjustments[0] = min(0.5, el['radius'] / max(min(el['w'], el['h']), 1))
    if el.get('fill'):
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(el['fill'])
    else:
        shp.fill.background()
    if el.get('stroke'):
        shp.line.color.rgb = rgb(el['stroke'])
        shp.line.width = Pt(el.get('lw', 0.8))
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ''


def add_line(slide, el):
    c = slide.shapes.add_connector(1, P(el['x1']), P(el['y1']), P(el['x2']), P(el['y2']))
    c.line.color.rgb = rgb(el['color'])
    c.line.width = Pt(el['lw'])
    if el.get('dash'):
        ln = c.line._get_or_add_ln()
        pd = ln.makeelement(qn('a:prstDash'), {'val': 'sysDot'})
        ln.append(pd)


def add_table(slide, el):
    rows, cols = len(el['rows']), len(el['cols'])
    total_h = sum(r['h'] for r in el['rows'])
    gf = slide.shapes.add_table(rows, cols, P(el['x']), P(el['y']), P(el['w']), P(total_h))
    tbl = gf.table
    # 기본 표 스타일 끄기
    tblpr = gf._element.graphic.graphicData.tbl.tblPr
    tblpr.set('firstRow', '0')
    tblpr.set('bandRow', '0')
    sid = tblpr.find(qn('a:tableStyleId'))
    if sid is not None:
        sid.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'   # No Style, Table Grid
    for j, w in enumerate(el['cols']):
        tbl.columns[j].width = P(w)
    pad = el['pad']
    for i, r in enumerate(el['rows']):
        tbl.rows[i].height = P(r['h'])
        for j, c in enumerate(r['cells']):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(c.get('fill') or LY.CELL)
            cell.margin_left = cell.margin_right = P(pad)
            cell.margin_top = cell.margin_bottom = P(5)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            tf = cell.text_frame
            _fill_tf(tf, c['lines'], c['size'], c.get('align', 'left'), c['size'] * 1.3)
            _cell_border(cell, LY.LINE)


def _cell_border(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        ln = tcpr.find(qn(tag))
        if ln is None:
            ln = tcpr.makeelement(qn(tag), {'w': '7620'})
            tcpr.append(ln)
        sf = ln.makeelement(qn('a:solidFill'), {})
        clr = sf.makeelement(qn('a:srgbClr'), {'val': color.lstrip('#')})
        sf.append(clr)
        ln.append(sf)


def render(d, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = P(LY.W), P(LY.H)
    blank = prs.slide_layouts[6]
    m = d.meta
    cp = prs.core_properties
    cp.title = f"{m['doc_no']} {m['title']}"
    cp.author = m['owner']
    cp.last_modified_by = m['owner']
    cp.subject = f"{m['access']} · {m.get('audience', '')}"
    cp.keywords = ', '.join(m.get('keywords', []))[:250]
    cp.category = m['access']
    for pg in d.pages:
        s = prs.slides.add_slide(blank)
        for el in pg.els:
            k = el['k']
            if k == 'rect':
                add_rect(s, el)
            elif k == 'line':
                add_line(s, el)
            elif k == 'text':
                add_text(s, el)
            elif k == 'image':
                s.shapes.add_picture(el['path'], P(el['x']), P(el['y']), P(el['w']), P(el['h']))
            elif k == 'table':
                add_table(s, el)
    prs.save(path)
